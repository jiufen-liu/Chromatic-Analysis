from dataclasses import replace
import xml.etree.ElementTree as ET

import pytest

from qtx_core.analysis import color_result
from qtx_core.colorimetry import reflectance_to_xyz_lab, reflectances_to_xyz_lab, xyz_to_lab
from qtx_core.cpx_io import parse_cpx_file, export_cpx_file
from qtx_core.models import Sample
from qtx_core.qtx_parser import parse_qtx_text, export_qtx_file


def cpx(tmp_path, illuminant="A", observer=2, spectrum=True, xyz=True, lab=True):
    curve = "".join(f'<Sp Wv="{w}">{20 + (w - 360)/20}</Sp>' for w in range(360, 701, 10))
    xml = f"""<SampleList><ViewingConditions><Illuminant>{illuminant}</Illuminant><Observer>d{observer}</Observer></ViewingConditions><Samples><Sample><Name>test</Name>
    {('<Spectrum>'+curve+'</Spectrum>') if spectrum else ''}
    {'<XYZ><X>20</X><Y>21</Y><Z>22</Z></XYZ>' if xyz else ''}
    {'<Lab><L>53</L><a>4</a><b>5</b></Lab>' if lab else ''}
    </Sample></Samples></SampleList>"""
    p=tmp_path/'test.cpx'; p.write_text(xml); return p


def test_cpx_alternate_condition_is_normalized_to_sample_contract(tmp_path):
    sm=parse_cpx_file(cpx(tmp_path)).samples[0]
    expected=reflectance_to_xyz_lab(sm.reflectance,'D65',sm.wavelengths,10)
    assert sm.xyz_d65_10 == pytest.approx(expected[0])
    assert color_result(sm,'D65').lab == pytest.approx(expected[1])
    assert sm.raw['CPX_VIEWING_ILLUMINANT']=='A'


def test_cpx_without_lab_derives_from_xyz(tmp_path):
    sm=parse_cpx_file(cpx(tmp_path,'D65',10,spectrum=False,lab=False)).samples[0]
    assert sm.lab_d65_10 == pytest.approx(xyz_to_lab((20,21,22),'D65',10))


def test_cpx_alternate_condition_without_spectrum_is_not_fabricated(tmp_path):
    with pytest.raises(ValueError,match='光谱'):
        parse_cpx_file(cpx(tmp_path,spectrum=False))


def test_cpx_export_cache_matches_declared_viewing_condition(tmp_path):
    sm=parse_cpx_file(cpx(tmp_path,'D65',10)).samples[0]
    out=tmp_path/'out.cpx'; export_cpx_file(out,'test',[sm],1,illuminant='A',observer=2)
    rt=ET.parse(out).getroot(); xyz=rt.find('./Samples/Sample/XYZ')
    actual=tuple(float(xyz.findtext(k)) for k in ('X','Y','Z'))
    assert actual == pytest.approx(color_result(sm,'A',2).xyz,abs=1e-5)
    assert parse_cpx_file(out).samples[0].xyz_d65_10 == pytest.approx(reflectance_to_xyz_lab(sm.reflectance,'D65',sm.wavelengths,10)[0])


@pytest.mark.parametrize('value',['10,broken,30','10,,30','10,nan,30','10,inf,30'])
def test_qtx_corrupt_curve_does_not_shift_or_poison_data(value):
    with pytest.raises(ValueError,match='光谱|反射率'):
        parse_qtx_text(f'[BATCH_DATA]\nBAT_NAME=broken\nBAT_R={value}\nBAT_REFLPOINTS=3\nBAT_X=20\nBAT_Y=20\nBAT_Z=20')


def test_qtx_declared_point_count_is_checked():
    with pytest.raises(ValueError,match='光谱|反射率'):
        parse_qtx_text('[BATCH_DATA]\nBAT_R=10,20,30\nBAT_REFLPOINTS=4\nBAT_X=20\nBAT_Y=20\nBAT_Z=20')


def test_qtx_transformed_target_export_does_not_restore_stale_raw(tmp_path):
    sm=Sample('id','target','STD',(20,21,22),xyz_to_lab((20,21,22)),(),raw={'STD_TX':'1','STD_TY':'2','STD_TZ':'3'})
    out=tmp_path/'out.qtx'; export_qtx_file(out,[sm])
    assert parse_qtx_text(out.read_text(encoding='utf-8-sig'))[0].xyz_d65_10 == pytest.approx(sm.xyz_d65_10)


@pytest.mark.parametrize('grid',[(360,370,370),(380,370,360),(360,float('nan'),380)])
@pytest.mark.parametrize('batch',[False,True])
def test_invalid_wavelength_grid_is_rejected(grid,batch):
    with pytest.raises(ValueError,match='波长'):
        if batch: reflectances_to_xyz_lab([[20,30,40],[30,40,50]],wavelengths=grid)
        else: reflectance_to_xyz_lab([20,30,40],wavelengths=grid)


def test_apparent_reflectance_above_100_remains_supported():
    xyz,lab=reflectance_to_xyz_lab([110]*35)
    assert xyz[1] == pytest.approx(110,abs=.1)
    assert lab[0]>100


def test_excel_preview_rgb_agrees_with_explicit_hex(tmp_path):
    from openpyxl import load_workbook
    from qtx_app.excel_exchange import export_client_card
    sm=Sample('id','test','STD',(20,21,22),(50,4,5),())
    out=tmp_path/'preview.xlsx'; export_client_card(str(out),'test',[sm],1,display_hexes=['#123456'])
    wb=load_workbook(out); ws=wb['颜色数据']
    assert tuple(ws.cell(2,col).value for col in (14,15,16)) == (18,52,86)
    assert ws.cell(2,17).value=='#123456'


def test_excel_internal_metadata_requires_explicit_opt_in(tmp_path):
    from openpyxl import load_workbook
    from qtx_app.excel_exchange import export_client_card
    sm=Sample('id','test','STD',(20,21,22),(50,4,5),(),raw={'ATTR_PrivateRecipe':'internal'})
    out=tmp_path/'internal.xlsx'; export_client_card(str(out),'test',[sm],include_internal_metadata=True)
    wb=load_workbook(out)
    assert wb['属性数据'].cell(2,3).value=='internal'


def test_munsell_runtime_dependency_is_available():
    from qtx_core.colorimetry import munsell_hue_order_from_xyz
    import math
    hue, notation, value, chroma = munsell_hue_order_from_xyz((20,21,10))
    assert math.isfinite(hue) and notation
    assert value > 0 and chroma > 0
