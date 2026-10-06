from dataclasses import asdict, replace
from pathlib import Path
from types import SimpleNamespace
import math

import pytest

from qtx_core.models import Sample
from qtx_app.main_window import MainWindow
from qtx_app import storage_v2


def sample():
    return Sample('same-id','测量样','STD',(20,21,22),(50,1,2),(40.0,)*35,source_file='same.qtx')


def owner():
    return SimpleNamespace(illuminant='D65',observer=10,_lab_cache={},_xyz_cache={},
        active_tool_index=lambda:0,_science_sample_signature=MainWindow._science_sample_signature)


def test_lab_cache_changes_when_measurement_changes_under_same_key():
    obj=owner(); original=sample(); edited=replace(original,lab_d65_10=(60,4,5))
    assert MainWindow.sample_lab(obj,original)==(50,1,2)
    assert MainWindow.sample_lab(obj,edited)==(60,4,5)


def test_xyz_cache_changes_when_measurement_changes_under_same_key():
    obj=owner(); original=sample(); edited=replace(original,xyz_d65_10=(30,31,32))
    assert MainWindow.sample_xyz(obj,original)==(20,21,22)
    assert MainWindow.sample_xyz(obj,edited)==(30,31,32)


def test_workbench_reloads_after_in_place_spectrum_or_lab_edit():
    obj=SimpleNamespace()
    obj._deserialize_sample=MainWindow._deserialize_sample.__get__(obj)
    payload=asdict(sample()); wb={'samples_data':[payload]}
    first=MainWindow._workbench_samples(obj,wb)
    payload['lab_d65_10']=[62,3,4]; payload['reflectance']=[60.0]*35
    edited=MainWindow._workbench_samples(obj,wb)
    assert edited[0].lab_d65_10==(62,3,4)
    assert edited[0].reflectance==(60.0,)*35
    assert MainWindow._workbench_samples(obj,wb) is edited
    assert first[0].lab_d65_10==(50,1,2)


@pytest.mark.parametrize('kind',['card','workbench'])
def test_mirror_fsync_failure_preserves_last_good_snapshot(tmp_path,monkeypatch,kind):
    monkeypatch.setattr(storage_v2,'color_card_root',lambda:tmp_path)
    monkeypatch.setattr(storage_v2,'color_card_snapshot_root',lambda *args,**kw:tmp_path)
    monkeypatch.setattr(storage_v2,'user_workbench_snapshot_root',lambda *args,**kw:tmp_path)
    if kind=='card':
        data={'card_id':'0123456789abcdef','name':'original','customer':'test'}
        write=lambda value:storage_v2.write_color_card_snapshot(value,1,'test')
    else:
        data={'workbench_id':'0123456789abcdef','name':'original'}
        write=lambda value:storage_v2.write_workbench_snapshot(value,1,'test')
    path=write(data); assert path and path.exists()
    old=path.read_bytes()
    def fail_sync(fd): raise OSError('simulated disk write failure')
    monkeypatch.setattr('os.fsync',fail_sync)
    assert write(dict(data,name='renamed')) is None
    assert path.exists() and path.read_bytes()==old
    assert not list(tmp_path.rglob('*.tmp'))


def test_shared_preview_matches_ui_excel_and_cpx_conventions():
    from qtx_core.preview import lab_to_srgb_preview
    from qtx_app.main_window import lab_to_qcolor, fluorescent_lab_to_qcolor
    from qtx_app.excel_exchange import _lab_to_hex, _lab_to_rgb
    from qtx_core.cpx_io import _rgb_from_lab
    for lab in [(0,0,0),(50,1,2),(80,100,-80),(110,-50,100)]:
        preview=lab_to_srgb_preview(lab)
        ui=lab_to_qcolor(lab)
        assert _lab_to_rgb(lab)==preview.rgb8
        assert _lab_to_hex(lab)==preview.hex[1:]
        assert all(abs(x-y)<=1 for x,y in zip((ui.red(),ui.green(),ui.blue()),preview.rgb8))
        clipped=lab_to_srgb_preview(lab,strategy='clip')
        assert _rgb_from_lab(lab)==clipped.rgb8
        assert fluorescent_lab_to_qcolor(lab).name().upper()==clipped.hex
    assert lab_to_srgb_preview((110,-50,100)).lightness_clamped
    assert lab_to_srgb_preview((80,100,-80)).gamut_mapped


@pytest.mark.parametrize('lab',[(math.nan,0,0),(50,math.inf,2),(50,1)])
def test_invalid_lab_cannot_be_exported_as_plausible_preview(lab):
    from qtx_core.preview import lab_to_srgb_preview
    with pytest.raises(ValueError): lab_to_srgb_preview(lab)


def test_failed_json_serialization_preserves_existing_workfile(tmp_path):
    from qtx_app.workfile import write_workfile, read_workfile
    path=tmp_path/'saved.chromatic'
    payload={'samples':[],'slots':[],'title':'original'}
    write_workfile(path,payload); original=path.read_bytes()
    with pytest.raises(ValueError):write_workfile(path,dict(payload,invalid=math.nan))
    assert path.read_bytes()==original
    assert read_workfile(path)['title']=='original'
    assert not list(tmp_path.glob('*.tmp'))


def test_concurrent_json_writes_leave_one_complete_document(tmp_path):
    from concurrent.futures import ThreadPoolExecutor
    from qtx_app.atomic_json import atomic_write_json
    import json
    path=tmp_path/'config.json'
    def write(i): atomic_write_json(path,{'sequence':i,'data':[i]*100})
    with ThreadPoolExecutor(max_workers=4) as pool:list(pool.map(write,range(12)))
    result=json.loads(path.read_text())
    assert result['data']==[result['sequence']]*100
    assert not list(tmp_path.glob('*.tmp'))


@pytest.mark.parametrize('kind',['card','workbench'])
def test_successful_mirror_rename_retires_old_copy_only_after_commit(tmp_path,monkeypatch,kind):
    monkeypatch.setattr(storage_v2,'color_card_root',lambda:tmp_path)
    monkeypatch.setattr(storage_v2,'color_card_snapshot_root',lambda *args,**kw:tmp_path)
    monkeypatch.setattr(storage_v2,'user_workbench_snapshot_root',lambda *args,**kw:tmp_path)
    if kind=='card':
        data={'card_id':'0123456789abcdef','name':'original','customer':'test'}
        write=lambda value:storage_v2.write_color_card_snapshot(value,1,'test')
    else:
        data={'workbench_id':'0123456789abcdef','name':'original'}
        write=lambda value:storage_v2.write_workbench_snapshot(value,1,'test')
    old=write(data); new=write(dict(data,name='renamed'))
    assert new.exists() and not old.exists()


def test_computed_results_keep_observer_and_commercial_spd_provenance():
    from qtx_core.analysis import color_result, analyse_pair
    source=sample()
    stored=color_result(source,'D65')
    assert stored.coordinate_source=='stored_d65_10' and stored.observer_degrees==10
    alternate=color_result(source,'U30',2)
    assert alternate.coordinate_source=='spectrum' and alternate.observer_degrees==2
    assert '兼容光谱' in alternate.illuminant_note
    pair=analyse_pair(source,source,'A',observer_degrees=2)
    assert pair.observer_degrees==2 and pair.reference_illuminant=='D65'


def test_condition_banner_distinguishes_commercial_approximation():
    from qtx_app.condition_status import condition_status
    text,tooltip,approximate=condition_status('D65',10)
    assert 'D65 / 10°' in text and '屏幕预览' in text and not approximate
    text,tooltip,approximate=condition_status('LEDT8G',2)
    assert '兼容光谱' in text and approximate and '客户' in tooltip


def test_library_fallback_is_recorded_instead_of_silently_relabelled():
    obj=owner(); obj.illuminant='A'
    incomplete=replace(sample(),reflectance=())
    assert MainWindow.sample_lab(obj,incomplete)==incomplete.lab_d65_10
    assert obj._preview_condition_failures
    assert '测量样' in next(iter(obj._preview_condition_failures.values()))


def test_excel_exports_treat_untrusted_name_and_attribute_as_text(tmp_path):
    from openpyxl import load_workbook
    from qtx_app.excel_exchange import export_client_card
    expression='=HYPERLINK("https://example.invalid", "name")'
    source=replace(sample(),display_name=expression,raw={'ATTR_Note':expression})
    target=tmp_path/'literal.xlsx'
    export_client_card(str(target),'test',[source],include_internal_metadata=True)
    wb=load_workbook(target,data_only=False)
    matches=[cell for sheet in wb for row in sheet for cell in row if cell.value==expression]
    assert len(matches)>=3 and all(cell.data_type=='s' for cell in matches)
    assert wb['颜色数据'].cell(2,6).data_type=='n'


@pytest.mark.parametrize('component',['lab','xyz'])
def test_nondefault_condition_hydrates_only_requested_lightweight_sample(component):
    from qtx_core.colorimetry import reflectance_to_xyz_lab
    full=sample()
    light=replace(full,reflectance=(),wavelengths=(),raw={'__P3_LIGHTWEIGHT__':'1'})
    obj=owner(); obj.illuminant='U30'; requested=[]
    obj._is_lightweight_library_sample=MainWindow._is_lightweight_library_sample
    def load(key): requested.append(key); return full
    obj.store=SimpleNamespace(load_sample_by_key=load)
    obj._full_library_sample=MainWindow._full_library_sample.__get__(obj)
    expected=reflectance_to_xyz_lab(full.reflectance,'U30',full.wavelengths,10)
    call=MainWindow.sample_lab if component=='lab' else MainWindow.sample_xyz
    actual=call(obj,light)
    assert actual==pytest.approx(expected[1 if component=='lab' else 0])
    assert call(obj,light)==actual and len(requested)==1
    assert not getattr(obj,'_preview_condition_failures',{})
