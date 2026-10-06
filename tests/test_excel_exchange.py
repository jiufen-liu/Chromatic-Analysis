from openpyxl import Workbook, load_workbook

from qtx_app.excel_exchange import export_client_card, import_mpc_data_workbook


def _mpc_workbook(path):
    wb = Workbook()
    cie = wb.active; cie.title = "CIE Coordinates"
    refl = wb.create_sheet("Reflectance Data")
    attrs = wb.create_sheet("Attributes")
    cie.append([]); cie.append([]); cie.append([])
    cie.append([None, "Sample", "X", "Y", "Z", "L", "a", "b", "C", "h"])
    cie.append([None, "TEST-1 sample", 20, 21, 22, 52, 1, -2, 2.24, 296.6])
    refl.append([]); refl.append([]); refl.append([])
    refl.append([None, "Sample", "Checksum", *[f"{w}nm" for w in range(360, 701, 10)]])
    refl.append([None, "TEST-1 sample", 0, *[40.0] * 35])
    attrs.append([]); attrs.append([]); attrs.append([])
    attrs.append([None, "Sample", "Ingredient 1", "Yarn Color Cost RMB/kg", "Commercial name", "Rank"])
    attrs.append([None, "TEST-1 sample", "SECRET", 9.5, "Test color", "1-1"])
    wb.save(path)


def test_mpc_excel_import_and_safe_client_export(tmp_path):
    source = tmp_path / "mpc.xlsx"
    _mpc_workbook(source)
    result = import_mpc_data_workbook(str(source))
    assert len(result.samples) == 1
    assert len(result.samples[0].reflectance) == 35
    assert result.layout[0]["column"] == 0
    assert result.layout[0]["row"] == 0
    assert result.internal_fields_present

    target = tmp_path / "client.xlsx"
    export_client_card(str(target), "Client card", [result.samples[0], None], 2)
    exported = load_workbook(target, data_only=True)
    values = [str(cell.value or "") for sheet in exported for row in sheet.iter_rows() for cell in row]
    joined = " ".join(values)
    assert "SECRET" not in joined
    assert "Yarn Color Cost" not in joined
