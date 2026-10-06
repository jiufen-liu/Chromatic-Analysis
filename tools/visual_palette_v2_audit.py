from __future__ import annotations

import argparse
import json
import math
import statistics
import sys
from datetime import datetime
from pathlib import Path
from time import perf_counter

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

from qtx_core import parse_qtx_file, visual_palette_layout_v2
from qtx_core.colorimetry import delta_e
from qtx_app.excel_exchange import _lab_to_hex


def _pct(vals, q: float) -> float:
    vals = sorted(float(v) for v in vals)
    if not vals:
        return 0.0
    idx = min(len(vals) - 1, max(0, int(round((len(vals) - 1) * q))))
    return vals[idx]


def _summary(vals) -> dict[str, float]:
    vals = [float(v) for v in vals]
    if not vals:
        return {"count": 0, "mean": 0.0, "median": 0.0, "p95": 0.0, "max": 0.0}
    return {
        "count": len(vals),
        "mean": statistics.fmean(vals),
        "median": statistics.median(vals),
        "p95": _pct(vals, 0.95),
        "max": max(vals),
    }


def _thin_border():
    side = Side(style="thin", color="D9E0EA")
    return Border(left=side, right=side, top=side, bottom=side)


def _safe_hex(lab) -> str:
    try:
        return _lab_to_hex(tuple(float(x) for x in lab)).replace("#", "")
    except Exception:
        return "D9D9D9"


def _short_hue_distance(a: float, b: float) -> float:
    d = abs(float(a) - float(b)) % 360.0
    return min(d, 360.0 - d)


def _row_metrics(grid, by_key, info):
    rows = []
    within_de = []
    within_hue = []
    within_L = []
    for r_idx, row in enumerate(grid):
        keys = [k for k in row if k is not None]
        if not keys:
            continue
        labs = [by_key[k].lab_d65_10 for k in keys]
        Ls = [float(x[0]) for x in labs]
        Cs = [math.hypot(float(x[1]), float(x[2])) for x in labs]
        hs = [math.degrees(math.atan2(float(x[2]), float(x[1]))) % 360.0 if c > 1e-12 else 0.0 for x, c in zip(labs, Cs)]
        des = []
        hds = []
        for i in range(len(keys) - 1):
            de = float(delta_e(labs[i], labs[i + 1], "CIE 2000"))
            des.append(de); within_de.append(de)
            if Cs[i] > 5.0 and Cs[i + 1] > 5.0:
                hd = _short_hue_distance(hs[i], hs[i + 1])
                hds.append(hd); within_hue.append(hd)
            within_L.append(abs(Ls[i] - Ls[i + 1]))
        fams = [info[k]["family"] for k in keys]
        rows.append({
            "row": r_idx + 1,
            "count": len(keys),
            "family": fams[0] if len(set(fams)) == 1 else " / ".join(dict.fromkeys(fams)),
            "L_min": min(Ls), "L_max": max(Ls), "L_span": max(Ls) - min(Ls),
            "C_min": min(Cs), "C_max": max(Cs),
            "de_mean": statistics.fmean(des) if des else 0.0,
            "de_max": max(des) if des else 0.0,
            "hue_jump_max": max(hds) if hds else 0.0,
        })
    return rows, _summary(within_de), _summary(within_hue), _summary(within_L)


def _write_preview_png(path: Path, grid, by_key, info, columns: int) -> bool:
    try:
        from PIL import Image, ImageDraw, ImageFont
    except Exception:
        return False

    cell_w, cell_h, sw_h = 220, 112, 66
    left_label = 120
    width = left_label + columns * cell_w
    height = max(1, len(grid)) * cell_h
    im = Image.new("RGB", (width, height), "white")
    draw = ImageDraw.Draw(im)

    font = None
    small = None
    for fp in (
        Path("C:/Windows/Fonts/msyh.ttc"),
        Path("C:/Windows/Fonts/msyh.ttf"),
        Path("C:/Windows/Fonts/simhei.ttf"),
    ):
        if fp.exists():
            try:
                font = ImageFont.truetype(str(fp), 14)
                small = ImageFont.truetype(str(fp), 11)
                break
            except Exception:
                pass
    if font is None:
        font = ImageFont.load_default(); small = font

    last_family = None
    rank = 0
    for r_idx, row in enumerate(grid):
        y = r_idx * cell_h
        keys = [k for k in row if k is not None]
        if not keys:
            draw.line((0, y + cell_h // 2, width, y + cell_h // 2), fill=(225, 230, 236), width=1)
            last_family = None
            continue
        family = info[keys[0]]["family"]
        if family != last_family:
            draw.text((6, y + 8), family, fill=(45, 55, 72), font=small)
            last_family = family
        for c_idx, key in enumerate(row):
            if key is None:
                continue
            rank += 1
            sm = by_key[key]
            lab = tuple(float(x) for x in sm.lab_d65_10)
            L, a, b = lab
            C = math.hypot(a, b)
            h = math.degrees(math.atan2(b, a)) % 360.0 if C > 1e-12 else 0.0
            hx = _safe_hex(lab)
            rgb = tuple(int(hx[i:i+2], 16) for i in (0, 2, 4))
            x = left_label + c_idx * cell_w
            draw.rounded_rectangle((x + 4, y + 4, x + cell_w - 8, y + sw_h), radius=8, fill=rgb, outline=(215, 220, 228))
            name = sm.display_name or "(unnamed)"
            if len(name) > 24:
                name = name[:23] + "…"
            draw.text((x + 7, y + sw_h + 5), f"{rank:03d} {name}", fill=(20, 25, 32), font=small)
            draw.text((x + 7, y + sw_h + 24), f"L {L:.1f}  C {C:.1f}  h {h:.0f}°", fill=(85, 95, 110), font=small)
    im.save(path)
    return True


def main() -> int:
    ap = argparse.ArgumentParser(description="Chromatic Analysis Visual Palette V2 offline audit")
    ap.add_argument("qtx", type=Path)
    ap.add_argument("--columns", type=int, default=6)
    args = ap.parse_args()

    qtx = args.qtx.expanduser().resolve()
    if not qtx.exists():
        print(f"文件不存在: {qtx}")
        return 2

    print("Chromatic Analysis · 视觉色卡编排 V2 离线验证")
    print("=" * 68)
    print(f"文件: {qtx}")
    print("说明：本工具只生成报告和完整预览，不修改正式色卡编排，不写入色库。")
    print("V2 采用二维色卡逻辑：色相分区 × 明度行 × 彩度/子色相列。")
    print()

    t0 = perf_counter()
    samples = parse_qtx_file(qtx)
    parse_s = perf_counter() - t0
    rows = []
    by_key = {}
    for i, sm in enumerate(samples, 1):
        key = f"S{i:04d}"
        by_key[key] = sm
        rows.append((key, sm.display_name or key, sm.lab_d65_10))

    t1 = perf_counter()
    result = visual_palette_layout_v2(rows, columns=args.columns, lightness_band=10.0, separator_rows=1)
    arrange_s = perf_counter() - t1
    grid = result["grid"]
    info = result["slot_info"]

    if len(result["keys"]) != len(samples):
        raise RuntimeError(f"完整性检查失败：输入 {len(samples)}，输出 {len(result['keys'])}")

    row_metrics, de_stats, hue_stats, L_stats = _row_metrics(grid, by_key, info)

    out_dir = ROOT / "diagnostics_reports"
    out_dir.mkdir(exist_ok=True)
    stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    xlsx_path = out_dir / f"VPA_V2_AUDIT_{stamp}.xlsx"
    txt_path = out_dir / f"VPA_V2_AUDIT_{stamp}.txt"
    json_path = out_dir / f"VPA_V2_AUDIT_{stamp}.json"
    png_path = out_dir / f"VPA_V2_PREVIEW_{stamp}.png"

    wb = Workbook()
    ws = wb.active
    ws.title = "总览"
    summary = [
        ("项目", "结果 / 定义"),
        ("输入文件", str(qtx)),
        ("输入色样", len(samples)),
        ("输出色样", len(result["keys"])),
        ("完整性", "PASS" if len(result["keys"]) == len(samples) else "FAIL"),
        ("列数", result["columns"]),
        ("总布局行数（含分隔行）", len(grid)),
        ("近中性色块", result["neutral_count"]),
        ("QTX解析时间(s)", round(parse_s, 4)),
        ("V2编排时间(s)", round(arrange_s, 4)),
        ("行内相邻 ΔE00 mean / p95 / max", f"{de_stats['mean']:.3f} / {de_stats['p95']:.3f} / {de_stats['max']:.3f}"),
        ("行内彩色相邻 hue jump mean / p95 / max", f"{hue_stats['mean']:.2f}° / {hue_stats['p95']:.2f}° / {hue_stats['max']:.2f}°"),
        ("行内相邻 |ΔL*| mean / p95 / max", f"{L_stats['mean']:.3f} / {L_stats['p95']:.3f} / {L_stats['max']:.3f}"),
        ("V2定位", "离线验证；不是标准色差公式；不是 Munsell 替代品；尚未进入正式色卡编排菜单。"),
        ("结构", "12个30°色相区；每区按10 L*明度切片；行内按15°子色相方向 + C*整理；白/灰/黑单独区。"),
    ]
    for r, (a, b) in enumerate(summary, 1):
        ws.cell(r, 1, a); ws.cell(r, 2, b)
        if r == 1:
            ws.cell(r, 1).font = ws.cell(r, 2).font = Font(bold=True)
            ws.cell(r, 1).fill = ws.cell(r, 2).fill = PatternFill("solid", fgColor="E8EEF7")
    ws.column_dimensions["A"].width = 30
    ws.column_dimensions["B"].width = 95
    ws.sheet_view.showGridLines = False

    # Matrix-style palette preview in Excel.
    pv = wb.create_sheet("V2二维编排")
    pv.sheet_view.showGridLines = False
    border = _thin_border()
    current_family = None
    out_row = 1
    for r_idx, row in enumerate(grid):
        keys = [k for k in row if k is not None]
        if not keys:
            out_row += 1
            current_family = None
            continue
        family = info[keys[0]]["family"]
        if family != current_family:
            pv.cell(out_row, 1, family)
            pv.cell(out_row, 1).font = Font(bold=True, color="334155")
            pv.merge_cells(start_row=out_row, start_column=1, end_row=out_row, end_column=args.columns)
            out_row += 1
            current_family = family
        for c_idx, key in enumerate(row, 1):
            if key is None:
                continue
            sm = by_key[key]
            lab = tuple(float(x) for x in sm.lab_d65_10)
            cell = pv.cell(out_row, c_idx, sm.display_name or key)
            cell.fill = PatternFill("solid", fgColor=_safe_hex(lab))
            cell.border = border
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.font = Font(size=9)
            cell.comment = None
        pv.row_dimensions[out_row].height = 48
        out_row += 1
    for c in range(1, args.columns + 1):
        pv.column_dimensions[get_column_letter(c)].width = 24

    # Detailed records ordered by real grid location.
    det = wb.create_sheet("V2数据明细")
    headers = ["顺序", "Grid Row", "Grid Col", "色块", "名称", "Family", "Neutral Kind", "L*", "a*", "b*", "C*", "h°", "Lightness Band"]
    for c, h in enumerate(headers, 1):
        det.cell(1, c, h).font = Font(bold=True)
        det.cell(1, c).fill = PatternFill("solid", fgColor="E8EEF7")
    det.freeze_panes = "A2"
    rank = 0
    for key in result["keys"]:
        rank += 1
        sm = by_key[key]
        L, a, b = map(float, sm.lab_d65_10)
        C = math.hypot(a, b)
        h = math.degrees(math.atan2(b, a)) % 360.0 if C > 1e-12 else 0.0
        si = info[key]
        vals = [rank, si["row"] + 1, si["column"] + 1, "", sm.display_name or key, si["family"], si["neutral_kind"], L, a, b, C, h, si["lightness_band"]]
        rr = rank + 1
        for cc, val in enumerate(vals, 1):
            det.cell(rr, cc, val).border = border
        det.cell(rr, 4).fill = PatternFill("solid", fgColor=_safe_hex((L, a, b)))
    widths = [10, 10, 10, 9, 34, 24, 16, 10, 10, 10, 10, 10, 15]
    for c, w in enumerate(widths, 1):
        det.column_dimensions[get_column_letter(c)].width = w

    rm = wb.create_sheet("行质量指标")
    rh = ["Grid Row", "数量", "Family", "L* min", "L* max", "L* span", "C* min", "C* max", "行内ΔE00 mean", "行内ΔE00 max", "行内hue jump max"]
    for c, h in enumerate(rh, 1):
        rm.cell(1, c, h).font = Font(bold=True)
        rm.cell(1, c).fill = PatternFill("solid", fgColor="E8EEF7")
    for i, rr in enumerate(row_metrics, 2):
        vals = [rr["row"], rr["count"], rr["family"], rr["L_min"], rr["L_max"], rr["L_span"], rr["C_min"], rr["C_max"], rr["de_mean"], rr["de_max"], rr["hue_jump_max"]]
        for c, v in enumerate(vals, 1):
            rm.cell(i, c, v).border = border
    for c in range(1, len(rh) + 1):
        rm.column_dimensions[get_column_letter(c)].width = 18 if c != 3 else 28

    wb.save(xlsx_path)
    made_png = _write_preview_png(png_path, grid, by_key, info, args.columns)

    payload = {
        "input": str(qtx),
        "sample_count": len(samples),
        "output_count": len(result["keys"]),
        "integrity_pass": len(result["keys"]) == len(samples),
        "columns": result["columns"],
        "grid_rows": len(grid),
        "neutral_count": result["neutral_count"],
        "family_counts": result["family_counts"],
        "parse_seconds": parse_s,
        "arrange_seconds": arrange_s,
        "within_row_de00": de_stats,
        "within_row_hue_jump": hue_stats,
        "within_row_abs_delta_L": L_stats,
        "png_created": made_png,
    }
    json_path.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    txt_lines = [
        "Chromatic Analysis · 视觉色卡编排 V2 离线验证",
        "=" * 68,
        f"输入: {qtx}",
        f"色样: {len(samples)} -> {len(result['keys'])}；完整性: {'PASS' if payload['integrity_pass'] else 'FAIL'}",
        f"V2编排: {arrange_s:.4f} s；QTX解析: {parse_s:.4f} s",
        f"近中性色: {result['neutral_count']}；总网格行: {len(grid)}；列数: {result['columns']}",
        f"行内 ΔE00 mean/p95/max: {de_stats['mean']:.3f} / {de_stats['p95']:.3f} / {de_stats['max']:.3f}",
        f"行内 hue jump mean/p95/max: {hue_stats['mean']:.2f} / {hue_stats['p95']:.2f} / {hue_stats['max']:.2f} deg",
        f"行内 |ΔL*| mean/p95/max: {L_stats['mean']:.3f} / {L_stats['p95']:.3f} / {L_stats['max']:.3f}",
        "",
        "Family counts:",
    ]
    for k, v in result["family_counts"].items():
        txt_lines.append(f"  {k}: {v}")
    txt_lines += ["", f"Excel: {xlsx_path}", f"Preview: {png_path if made_png else 'Pillow unavailable'}", f"JSON: {json_path}"]
    txt_path.write_text("\n".join(txt_lines), encoding="utf-8")

    print("完成。")
    print(f"  Excel: {xlsx_path}")
    if made_png:
        print(f"  完整预览 PNG: {png_path}")
    print(f"  摘要: {txt_path}")
    print("请优先把 PNG + Excel 发给 ChatGPT；先验证视觉结构，再决定是否进入正式色卡编排。")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
