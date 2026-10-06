from __future__ import annotations

import argparse
import json
import math
import statistics
import sys
from datetime import datetime
from pathlib import Path
from time import perf_counter

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

from qtx_core import parse_qtx_file, visual_palette_layout_v2, visual_palette_layout_v21
from qtx_core.colorimetry import delta_e
from qtx_app.excel_exchange import _lab_to_hex


def _pct(vals, q: float) -> float:
    vals = sorted(float(v) for v in vals)
    if not vals:
        return 0.0
    idx = min(len(vals) - 1, max(0, int(round((len(vals) - 1) * q))))
    return vals[idx]


def _summary(vals) -> dict[str, float]:
    vals = [float(v) for v in vals]
    if not vals:
        return {"count": 0, "mean": 0.0, "median": 0.0, "p95": 0.0, "max": 0.0}
    return {
        "count": len(vals),
        "mean": statistics.fmean(vals),
        "median": statistics.median(vals),
        "p95": _pct(vals, 0.95),
        "max": max(vals),
    }


def _thin_border():
    side = Side(style="thin", color="D9E0EA")
    return Border(left=side, right=side, top=side, bottom=side)


def _safe_hex(lab) -> str:
    try:
        return _lab_to_hex(tuple(float(x) for x in lab)).replace("#", "")
    except Exception:
        return "D9D9D9"


def _short_hue_distance(a: float, b: float) -> float:
    d = abs(float(a) - float(b)) % 360.0
    return min(d, 360.0 - d)


def _row_metrics(grid, by_key, info):
    rows = []
    within_de = []
    within_hue = []
    within_L = []
    for r_idx, row in enumerate(grid):
        keys = [k for k in row if k is not None]
        if not keys:
            continue
        labs = [by_key[k].lab_d65_10 for k in keys]
        Ls = [float(x[0]) for x in labs]
        Cs = [math.hypot(float(x[1]), float(x[2])) for x in labs]
        hs = [math.degrees(math.atan2(float(x[2]), float(x[1]))) % 360.0 if c > 1e-12 else 0.0 for x, c in zip(labs, Cs)]
        des, hds = [], []
        for i in range(len(keys) - 1):
            de = float(delta_e(labs[i], labs[i + 1], "CIE 2000"))
            des.append(de); within_de.append(de)
            # Hue angle is not a valid quality signal for explicit neutral groups.
            fam = info[keys[i]]["family"]
            fam2 = info[keys[i+1]]["family"]
            if fam not in {"White", "Neutral Grey", "Black", "Neutral / White / Black"} and fam2 == fam and Cs[i] >= 10.0 and Cs[i + 1] >= 10.0:
                hd = _short_hue_distance(hs[i], hs[i + 1])
                hds.append(hd); within_hue.append(hd)
            within_L.append(abs(Ls[i] - Ls[i + 1]))
        fams = [info[k]["family"] for k in keys]
        rows.append({
            "row": r_idx + 1,
            "count": len(keys),
            "family": fams[0] if len(set(fams)) == 1 else " / ".join(dict.fromkeys(fams)),
            "L_min": min(Ls), "L_max": max(Ls), "L_span": max(Ls) - min(Ls),
            "C_min": min(Cs), "C_max": max(Cs),
            "de_mean": statistics.fmean(des) if des else 0.0,
            "de_max": max(des) if des else 0.0,
            "hue_jump_max": max(hds) if hds else 0.0,
        })
    return rows, _summary(within_de), _summary(within_hue), _summary(within_L)


def _section_transition_metrics(grid, by_key, info):
    """Measure row-to-row transitions only inside the same semantic family."""
    nonempty = [(idx, [k for k in row if k is not None]) for idx, row in enumerate(grid) if any(k is not None for k in row)]
    vals = []
    detail = []
    for (r1, a), (r2, b) in zip(nonempty, nonempty[1:]):
        if not a or not b:
            continue
        if info[a[0]]["family"] != info[b[0]]["family"]:
            continue
        k1, k2 = a[-1], b[0]
        de = float(delta_e(by_key[k1].lab_d65_10, by_key[k2].lab_d65_10, "CIE 2000"))
        vals.append(de)
        detail.append((r1 + 1, r2 + 1, info[k1]["family"], k1, k2, de))
    return _summary(vals), detail


def _boundary_conflicts(keys, by_key, info, threshold=3.0):
    out = []
    for i, ka in enumerate(keys):
        for kb in keys[i+1:]:
            fa, fb = info[ka]["family"], info[kb]["family"]
            if fa == fb:
                continue
            de = float(delta_e(by_key[ka].lab_d65_10, by_key[kb].lab_d65_10, "CIE 2000"))
            if de < threshold:
                out.append((de, ka, kb, fa, fb))
    out.sort(key=lambda x: x[0])
    return out


def _write_preview_png(path: Path, grid, by_key, info, columns: int) -> bool:
    try:
        from PIL import Image, ImageDraw, ImageFont
    except Exception:
        return False
    cell_w, cell_h, sw_h = 220, 112, 66
    left_label = 120
    width = left_label + columns * cell_w
    height = max(1, len(grid)) * cell_h
    im = Image.new("RGB", (width, height), "white")
    draw = ImageDraw.Draw(im)
    font = small = None
    for fp in (Path("C:/Windows/Fonts/msyh.ttc"), Path("C:/Windows/Fonts/msyh.ttf"), Path("C:/Windows/Fonts/simhei.ttf")):
        if fp.exists():
            try:
                font = ImageFont.truetype(str(fp), 14); small = ImageFont.truetype(str(fp), 11); break
            except Exception:
                pass
    if font is None:
        font = ImageFont.load_default(); small = font
    last_family = None
    rank = 0
    for r_idx, row in enumerate(grid):
        y = r_idx * cell_h
        keys = [k for k in row if k is not None]
        if not keys:
            continue
        family = info[keys[0]]["family"]
        if family != last_family:
            draw.line((0, y, width, y), fill=(180, 190, 204), width=2)
            draw.text((6, y + 8), family, fill=(45, 55, 72), font=small)
            last_family = family
        for c_idx, key in enumerate(row):
            if key is None:
                continue
            rank += 1
            sm = by_key[key]
            lab = tuple(float(x) for x in sm.lab_d65_10)
            L, a, b = lab
            C = math.hypot(a, b)
            h = math.degrees(math.atan2(b, a)) % 360.0 if C > 1e-12 else 0.0
            hx = _safe_hex(lab)
            rgb = tuple(int(hx[i:i+2], 16) for i in (0, 2, 4))
            x = left_label + c_idx * cell_w
            draw.rounded_rectangle((x + 4, y + 4, x + cell_w - 8, y + sw_h), radius=8, fill=rgb, outline=(215, 220, 228))
            name = sm.display_name or "(unnamed)"
            if len(name) > 24: name = name[:23] + "…"
            draw.text((x + 7, y + sw_h + 5), f"{rank:03d} {name}", fill=(20, 25, 32), font=small)
            draw.text((x + 7, y + sw_h + 24), f"L {L:.1f}  C {C:.1f}  h {h:.0f}°", fill=(85, 95, 110), font=small)
    im.save(path)
    return True


def _write_palette_sheet(wb, title, grid, by_key, info, columns):
    ws = wb.create_sheet(title)
    ws.sheet_view.showGridLines = False
    border = _thin_border()
    last_family = None
    out_row = 1
    for row in grid:
        keys = [k for k in row if k is not None]
        if not keys:
            continue
        family = info[keys[0]]["family"]
        if family != last_family:
            ws.cell(out_row, 1, family)
            ws.cell(out_row, 1).font = Font(bold=True, color="334155")
            ws.merge_cells(start_row=out_row, start_column=1, end_row=out_row, end_column=columns)
            out_row += 1
            last_family = family
        for c_idx, key in enumerate(row, 1):
            if key is None: continue
            sm = by_key[key]
            lab = tuple(float(x) for x in sm.lab_d65_10)
            cell = ws.cell(out_row, c_idx, sm.display_name or key)
            cell.fill = PatternFill("solid", fgColor=_safe_hex(lab))
            cell.border = border
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.font = Font(size=9)
        ws.row_dimensions[out_row].height = 48
        out_row += 1
    for c in range(1, columns + 1):
        ws.column_dimensions[get_column_letter(c)].width = 24
    return ws


def main() -> int:
    ap = argparse.ArgumentParser(description="Chromatic Analysis Visual Palette V2.1 offline audit")
    ap.add_argument("qtx", type=Path)
    ap.add_argument("--columns", type=int, default=6)
    args = ap.parse_args()
    qtx = args.qtx.expanduser().resolve()
    if not qtx.exists():
        print(f"文件不存在: {qtx}"); return 2

    print("Chromatic Analysis · 视觉色卡编排 V2.1 离线验证")
    print("=" * 68)
    print(f"文件: {qtx}")
    print("说明：同时计算 HF77 V2 与 HF78 V2.1，只生成报告/预览，不修改正式色卡。")
    print()

    t0 = perf_counter(); samples = parse_qtx_file(qtx); parse_s = perf_counter() - t0
    rows, by_key = [], {}
    for i, sm in enumerate(samples, 1):
        key = f"S{i:04d}"; by_key[key] = sm; rows.append((key, sm.display_name or key, sm.lab_d65_10))

    t1 = perf_counter(); old = visual_palette_layout_v2(rows, columns=args.columns, lightness_band=10.0, separator_rows=0); old_s = perf_counter()-t1
    t2 = perf_counter(); new = visual_palette_layout_v21(rows, columns=args.columns, lightness_span=7.5, separator_rows=0); new_s = perf_counter()-t2
    if len(new["keys"]) != len(samples):
        raise RuntimeError(f"完整性检查失败：输入 {len(samples)}，输出 {len(new['keys'])}")

    old_rows, old_de, old_hue, old_L = _row_metrics(old["grid"], by_key, old["slot_info"])
    new_rows, new_de, new_hue, new_L = _row_metrics(new["grid"], by_key, new["slot_info"])
    old_trans, _ = _section_transition_metrics(old["grid"], by_key, old["slot_info"])
    new_trans, trans_detail = _section_transition_metrics(new["grid"], by_key, new["slot_info"])
    conflicts = _boundary_conflicts(new["keys"], by_key, new["slot_info"], 3.0)
    low_conf = [k for k in new["keys"] if new["slot_info"][k].get("hue_confidence") == "LOW"]

    out_dir = ROOT / "diagnostics_reports"; out_dir.mkdir(exist_ok=True)
    stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    xlsx_path = out_dir / f"VPA_V21_AUDIT_{stamp}.xlsx"
    png_path = out_dir / f"VPA_V21_PREVIEW_{stamp}.png"
    txt_path = out_dir / f"VPA_V21_AUDIT_{stamp}.txt"
    json_path = out_dir / f"VPA_V21_AUDIT_{stamp}.json"

    _write_preview_png(png_path, new["grid"], by_key, new["slot_info"], args.columns)

    wb = Workbook(); ws = wb.active; ws.title = "总览"
    summary = [
        ("项目", "结果 / 定义"),
        ("输入文件", str(qtx)), ("输入色样", len(samples)), ("输出色样", len(new["keys"])), ("完整性", "PASS"),
        ("QTX解析时间(s)", round(parse_s,4)), ("V2时间(s)", round(old_s,4)), ("V2.1时间(s)", round(new_s,4)),
        ("V2 行内 ΔE00 mean / p95 / max", f"{old_de['mean']:.3f} / {old_de['p95']:.3f} / {old_de['max']:.3f}"),
        ("V2.1 行内 ΔE00 mean / p95 / max", f"{new_de['mean']:.3f} / {new_de['p95']:.3f} / {new_de['max']:.3f}"),
        ("V2.1 行内彩色 hue jump mean / p95 / max", f"{new_hue['mean']:.2f}° / {new_hue['p95']:.2f}° / {new_hue['max']:.2f}°"),
        ("V2.1 行内 |ΔL*| mean / p95 / max", f"{new_L['mean']:.3f} / {new_L['p95']:.3f} / {new_L['max']:.3f}"),
        ("V2 行间同Family ΔE00 mean / p95 / max", f"{old_trans['mean']:.3f} / {old_trans['p95']:.3f} / {old_trans['max']:.3f}"),
        ("V2.1 行间同Family ΔE00 mean / p95 / max", f"{new_trans['mean']:.3f} / {new_trans['p95']:.3f} / {new_trans['max']:.3f}"),
        ("V2.1 跨Family但 ΔE00<3 的边界冲突", len(conflicts)),
        ("V2.1 低Hue置信度彩色样(C*<10)", len(low_conf)),
        ("V2.1定位", "离线验证；保持12色相骨架；自适应明度层；白/灰/黑分区；仍未进入正式色卡编排。"),
    ]
    for r,(a,b) in enumerate(summary,1):
        ws.cell(r,1,a); ws.cell(r,2,b)
        if r==1:
            ws.cell(r,1).font=ws.cell(r,2).font=Font(bold=True)
            ws.cell(r,1).fill=ws.cell(r,2).fill=PatternFill("solid",fgColor="E8EEF7")
    ws.column_dimensions["A"].width=38; ws.column_dimensions["B"].width=90; ws.sheet_view.showGridLines=False

    _write_palette_sheet(wb, "V2.1二维编排", new["grid"], by_key, new["slot_info"], args.columns)

    det = wb.create_sheet("V2.1数据明细")
    headers=["顺序","Grid Row","Grid Col","色块","名称","Family","Neutral Kind","Hue Confidence","L*","a*","b*","C*","h°","Shelf"]
    for c,h in enumerate(headers,1): det.cell(1,c,h).font=Font(bold=True); det.cell(1,c).fill=PatternFill("solid",fgColor="E8EEF7")
    rank=0
    for r_idx,row in enumerate(new["grid"]):
        for c_idx,key in enumerate(row):
            if key is None: continue
            rank+=1; sm=by_key[key]; lab=tuple(float(x) for x in sm.lab_d65_10); L,a,b=lab; C=math.hypot(a,b); h=math.degrees(math.atan2(b,a))%360 if C>1e-12 else 0.0; inf=new["slot_info"][key]
            vals=[rank,r_idx+1,c_idx+1,"",sm.display_name or key,inf["family"],inf.get("neutral_kind",""),inf.get("hue_confidence",""),L,a,b,C,h,inf.get("lightness_band","")]
            for col,v in enumerate(vals,1): det.cell(rank+1,col,v)
            det.cell(rank+1,4).fill=PatternFill("solid",fgColor=_safe_hex(lab))
    det.freeze_panes="A2"; det.auto_filter.ref=f"A1:N{rank+1}"
    for col in range(1,15): det.column_dimensions[get_column_letter(col)].width=14
    det.column_dimensions["E"].width=45; det.column_dimensions["F"].width=20; det.column_dimensions["H"].width=18

    qm=wb.create_sheet("行质量指标")
    qheaders=["Grid Row","数量","Family","L* min","L* max","L* span","C* min","C* max","行内ΔE00 mean","行内ΔE00 max","行内hue jump max"]
    for c,h in enumerate(qheaders,1): qm.cell(1,c,h).font=Font(bold=True); qm.cell(1,c).fill=PatternFill("solid",fgColor="E8EEF7")
    for rr,row in enumerate(new_rows,2):
        vals=[row["row"],row["count"],row["family"],row["L_min"],row["L_max"],row["L_span"],row["C_min"],row["C_max"],row["de_mean"],row["de_max"],row["hue_jump_max"]]
        for c,v in enumerate(vals,1): qm.cell(rr,c,v)
    qm.freeze_panes="A2"; qm.auto_filter.ref=f"A1:K{len(new_rows)+1}"

    bc=wb.create_sheet("边界冲突")
    bheaders=["ΔE00","样本A","Family A","L*A","C*A","h°A","样本B","Family B","L*B","C*B","h°B","说明"]
    for c,h in enumerate(bheaders,1): bc.cell(1,c,h).font=Font(bold=True); bc.cell(1,c).fill=PatternFill("solid",fgColor="FDECEC")
    for r,(de,ka,kb,fa,fb) in enumerate(conflicts,2):
        sa,sb=by_key[ka],by_key[kb]; La,aa,ba=map(float,sa.lab_d65_10); Lb,ab,bb=map(float,sb.lab_d65_10); Ca=math.hypot(aa,ba); Cb=math.hypot(ab,bb); ha=math.degrees(math.atan2(ba,aa))%360 if Ca else 0; hb=math.degrees(math.atan2(bb,ab))%360 if Cb else 0
        vals=[de,sa.display_name or ka,fa,La,Ca,ha,sb.display_name or kb,fb,Lb,Cb,hb,"视觉很接近但被分到不同Family，优先人工复核分类边界"]
        for c,v in enumerate(vals,1): bc.cell(r,c,v)
    bc.freeze_panes="A2"; bc.auto_filter.ref=f"A1:L{max(2,len(conflicts)+1)}"; bc.column_dimensions["B"].width=42; bc.column_dimensions["G"].width=42; bc.column_dimensions["L"].width=50

    wb.save(xlsx_path)

    payload={
        "schema":1,"version":"V2.1","input":str(qtx),"sample_count":len(samples),"parse_s":parse_s,
        "v2":{"arrange_s":old_s,"row_de":old_de,"row_hue":old_hue,"row_L":old_L,"same_family_row_transition":old_trans},
        "v21":{"arrange_s":new_s,"row_de":new_de,"row_hue":new_hue,"row_L":new_L,"same_family_row_transition":new_trans,"boundary_conflicts_de_lt_3":len(conflicts),"low_hue_confidence":len(low_conf),"family_counts":new["family_counts"]},
    }
    json_path.write_text(json.dumps(payload,ensure_ascii=False,indent=2),encoding="utf-8")
    lines=["Chromatic Analysis · 视觉色卡编排 V2.1 离线验证","="*68,f"文件: {qtx}",f"输入/输出: {len(samples)} / {len(new['keys'])}  PASS",f"V2.1 行内ΔE00: mean={new_de['mean']:.3f}, p95={new_de['p95']:.3f}, max={new_de['max']:.3f}",f"V2.1 同Family行间ΔE00: mean={new_trans['mean']:.3f}, p95={new_trans['p95']:.3f}, max={new_trans['max']:.3f}",f"跨Family但ΔE00<3: {len(conflicts)}",f"低Hue置信度彩色样(C*<10): {len(low_conf)}","",f"预览: {png_path.name}",f"Excel: {xlsx_path.name}"]
    txt_path.write_text("\n".join(lines),encoding="utf-8")
    print("完成。")
    print(f"预览: {png_path}\nExcel: {xlsx_path}\nTXT: {txt_path}\nJSON: {json_path}")
    print("请把 PNG + XLSX 发给 ChatGPT。")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
