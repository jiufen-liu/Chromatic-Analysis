from __future__ import annotations

import argparse
import json
import math
import statistics
import sys
from collections import Counter
from datetime import datetime
from pathlib import Path
from time import perf_counter

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

from qtx_core import parse_qtx_file, visual_palette_layout_v24, visual_palette_layout_v241
from qtx_core.colorimetry import delta_e
from qtx_app.excel_exchange import _lab_to_hex


def _pct(vals, q: float) -> float:
    vals = sorted(float(v) for v in vals)
    if not vals:
        return 0.0
    idx = min(len(vals) - 1, max(0, int(round((len(vals) - 1) * q))))
    return vals[idx]


def _summary(vals) -> dict[str, float]:
    vals = [float(v) for v in vals]
    if not vals:
        return {"count": 0, "mean": 0.0, "median": 0.0, "p95": 0.0, "max": 0.0}
    return {
        "count": len(vals), "mean": statistics.fmean(vals), "median": statistics.median(vals),
        "p95": _pct(vals, 0.95), "max": max(vals),
    }


def _thin_border():
    side = Side(style="thin", color="D9E0EA")
    return Border(left=side, right=side, top=side, bottom=side)


def _safe_hex(lab) -> str:
    try:
        return _lab_to_hex(tuple(float(x) for x in lab)).replace("#", "")
    except Exception:
        return "D9D9D9"


def _short_hue_distance(a: float, b: float) -> float:
    d = abs(float(a) - float(b)) % 360.0
    return min(d, 360.0 - d)


def _family_adjacent(i: int, j: int) -> bool:
    if i < 0 or j < 0 or i >= 12 or j >= 12:
        return False
    d = abs(i - j)
    return d == 1 or d == 11


def _row_metrics(grid, by_key, info):
    rows = []; within_de=[]; within_hue=[]; within_L=[]; within_C=[]; c_monotonic_violations=0
    for r_idx,row in enumerate(grid):
        keys=[k for k in row if k is not None]
        if not keys: continue
        labs=[by_key[k].lab_d65_10 for k in keys]
        Ls=[float(x[0]) for x in labs]; Cs=[math.hypot(float(x[1]),float(x[2])) for x in labs]
        hs=[math.degrees(math.atan2(float(x[2]),float(x[1])))%360.0 if c>1e-12 else 0.0 for x,c in zip(labs,Cs)]
        des=[]; hds=[]
        for i in range(len(keys)-1):
            de=float(delta_e(labs[i],labs[i+1],"CIE 2000")); des.append(de); within_de.append(de)
            within_L.append(abs(Ls[i]-Ls[i+1])); within_C.append(abs(Cs[i]-Cs[i+1]))
            fam=info[keys[i]]["family"]; fam2=info[keys[i+1]]["family"]
            if fam==fam2 and fam not in {"Neutral Axis","Neutral Field"} and Cs[i]>=10.0 and Cs[i+1]>=10.0:
                hd=_short_hue_distance(hs[i],hs[i+1]); hds.append(hd); within_hue.append(hd)
            if fam==fam2 and fam not in {"Neutral Axis","Neutral Field"} and Cs[i+1]+1e-9<Cs[i]:
                c_monotonic_violations += 1
        fams=[info[k]["family"] for k in keys]
        rows.append({
            "row":r_idx+1,"count":len(keys),"family":fams[0] if len(set(fams))==1 else " / ".join(dict.fromkeys(fams)),
            "L_min":min(Ls),"L_max":max(Ls),"L_span":max(Ls)-min(Ls),
            "C_min":min(Cs),"C_max":max(Cs),"C_span":max(Cs)-min(Cs),
            "de_mean":statistics.fmean(des) if des else 0.0,"de_max":max(des) if des else 0.0,
            "hue_jump_max":max(hds) if hds else 0.0,
        })
    return {"rows":rows,"de":_summary(within_de),"hue":_summary(within_hue),"L":_summary(within_L),"C_step":_summary(within_C),
            "c_monotonic_violations":c_monotonic_violations,"row_c_span":_summary([r["C_span"] for r in rows])}


def _vertical_nearest_c_metrics(grid, by_key, info):
    """2-D atlas metric: adjacent L* rows, matched by nearest C* rather than row-end -> row-start."""
    nonempty=[(idx,[k for k in row if k is not None]) for idx,row in enumerate(grid) if any(k is not None for k in row)]
    vals=[]; detail=[]
    for (r1,a),(r2,b) in zip(nonempty,nonempty[1:]):
        if not a or not b: continue
        fam=info[a[0]]["family"]
        if fam != info[b[0]]["family"]: continue
        def C(k):
            L,aa,bb=(float(x) for x in by_key[k].lab_d65_10); return math.hypot(aa,bb)
        pairs=set()
        for ka in a:
            kb=min(b,key=lambda x:abs(C(x)-C(ka))); pairs.add((ka,kb))
        for kb in b:
            ka=min(a,key=lambda x:abs(C(x)-C(kb))); pairs.add((ka,kb))
        for ka,kb in pairs:
            de=float(delta_e(by_key[ka].lab_d65_10,by_key[kb].lab_d65_10,"CIE 2000"))
            vals.append(de); detail.append((r1+1,r2+1,fam,ka,kb,de))
    return _summary(vals), detail


def _cross_family_near_pairs_v241(keys, by_key, info, threshold=3.0):
    hue_bridges=[]; neutral_bridges=[]; conflicts=[]
    hue_samples=set(); neutral_samples=set(); conflict_samples=set()
    for i,ka in enumerate(keys):
        ia=int(info[ka].get("family_index",99)); fa=info[ka]["family"]
        for kb in keys[i+1:]:
            fb=info[kb]["family"]
            if fa==fb: continue
            de=float(delta_e(by_key[ka].lab_d65_10,by_key[kb].lab_d65_10,"CIE 2000"))
            if de>=threshold: continue
            ib=int(info[kb].get("family_index",99)); item=(de,ka,kb,fa,fb,ia,ib)
            if _family_adjacent(ia,ib):
                hue_bridges.append(item); hue_samples.update((ka,kb)); continue
            if ia==99 or ib==99:
                other = kb if ia==99 else ka
                other_info=info[other]
                # Near-neutral family transitions are expected when the other
                # swatch has weak/uncertain chroma.  Treat them as bridges,
                # not semantic failures.
                if other_info.get("hue_confidence") in {"VERY_LOW","LOW"} or float(other_info.get("C",99.0)) < 12.0:
                    neutral_bridges.append(item); neutral_samples.update((ka,kb)); continue
            conflicts.append(item); conflict_samples.update((ka,kb))
    hue_bridges.sort(key=lambda x:x[0]); neutral_bridges.sort(key=lambda x:x[0]); conflicts.sort(key=lambda x:x[0])
    return hue_bridges, neutral_bridges, conflicts, hue_samples, neutral_samples, conflict_samples


def _write_preview_png(path: Path, grid, by_key, info, columns: int) -> bool:
    try:
        from PIL import Image, ImageDraw, ImageFont
    except Exception:
        return False
    cell_w,cell_h,sw_h=220,112,66; left_label=120
    width=left_label+columns*cell_w; height=max(1,len(grid))*cell_h
    im=Image.new("RGB",(width,height),"white"); draw=ImageDraw.Draw(im); font=small=None
    for fp in (Path("C:/Windows/Fonts/msyh.ttc"),Path("C:/Windows/Fonts/msyh.ttf"),Path("C:/Windows/Fonts/simhei.ttf")):
        if fp.exists():
            try: font=ImageFont.truetype(str(fp),14); small=ImageFont.truetype(str(fp),11); break
            except Exception: pass
    if font is None: font=ImageFont.load_default(); small=font
    last_family=None; rank=0
    for r_idx,row in enumerate(grid):
        y=r_idx*cell_h; keys=[k for k in row if k is not None]
        if not keys: continue
        family=info[keys[0]]["family"]
        if family!=last_family:
            draw.line((0,y,width,y),fill=(180,190,204),width=2); draw.text((6,y+8),family,fill=(45,55,72),font=small); last_family=family
        for c_idx,key in enumerate(row):
            if key is None: continue
            rank+=1; sm=by_key[key]; lab=tuple(float(x) for x in sm.lab_d65_10); L,a,b=lab; C=math.hypot(a,b); h=math.degrees(math.atan2(b,a))%360.0 if C>1e-12 else 0.0
            hx=_safe_hex(lab); rgb=tuple(int(hx[i:i+2],16) for i in (0,2,4)); x=left_label+c_idx*cell_w
            draw.rounded_rectangle((x+4,y+4,x+cell_w-8,y+sw_h),radius=8,fill=rgb,outline=(215,220,228))
            name=sm.display_name or "(unnamed)"; name=name if len(name)<=24 else name[:23]+"…"; mark=" ◇" if info[key].get("boundary_zone") else ""
            draw.text((x+7,y+sw_h+5),f"{rank:03d} {name}{mark}",fill=(20,25,32),font=small)
            draw.text((x+7,y+sw_h+24),f"L {L:.1f}  C {C:.1f}  h {h:.0f}°",fill=(85,95,110),font=small)
    im.save(path); return True


def _write_palette_sheet(wb, title, grid, by_key, info, columns):
    ws=wb.create_sheet(title); ws.sheet_view.showGridLines=False; border=_thin_border(); last_family=None; out_row=1
    for row in grid:
        keys=[k for k in row if k is not None]
        if not keys: continue
        family=info[keys[0]]["family"]
        if family!=last_family:
            ws.cell(out_row,1,family); ws.cell(out_row,1).font=Font(bold=True,color="334155")
            ws.merge_cells(start_row=out_row,start_column=1,end_row=out_row,end_column=columns); out_row+=1; last_family=family
        for c_idx,key in enumerate(row,1):
            if key is None: continue
            sm=by_key[key]; lab=tuple(float(x) for x in sm.lab_d65_10); label=sm.display_name or key
            if info[key].get("boundary_zone"): label += "  ◇边界"
            cell=ws.cell(out_row,c_idx,label); cell.fill=PatternFill("solid",fgColor=_safe_hex(lab)); cell.border=border
            cell.alignment=Alignment(horizontal="center",vertical="center",wrap_text=True); cell.font=Font(size=9)
        ws.row_dimensions[out_row].height=48; out_row+=1
    for c in range(1,columns+1): ws.column_dimensions[get_column_letter(c)].width=24
    return ws


def _write_pairs(wb, title, pairs, by_key, fill, note):
    sh=wb.create_sheet(title); headers=["ΔE00","样本A","Family A","样本B","Family B","说明"]
    for c,h in enumerate(headers,1): sh.cell(1,c,h).font=Font(bold=True); sh.cell(1,c).fill=PatternFill("solid",fgColor=fill)
    for r,(de,ka,kb,fa,fb,ia,ib) in enumerate(pairs,2):
        vals=[de,by_key[ka].display_name or ka,fa,by_key[kb].display_name or kb,fb,note]
        for c,v in enumerate(vals,1): sh.cell(r,c,v)
    sh.freeze_panes="A2"; sh.auto_filter.ref=f"A1:F{max(2,len(pairs)+1)}"; sh.column_dimensions["B"].width=45; sh.column_dimensions["D"].width=45; sh.column_dimensions["F"].width=62


def _audit_one(qtx: Path, out_dir: Path, *, columns: int=6, prefix: str="VPA_V241") -> dict:
    t0=perf_counter(); samples=parse_qtx_file(qtx); parse_s=perf_counter()-t0
    rows=[]; by_key={}
    for i,sm in enumerate(samples,1):
        key=f"S{i:04d}"; by_key[key]=sm; rows.append((key,sm.display_name or key,sm.lab_d65_10))
    t1=perf_counter(); old=visual_palette_layout_v24(rows,columns=columns,lightness_span=6.0,soft_boundary_deg=4.0,separator_rows=0); old_s=perf_counter()-t1
    t2=perf_counter(); new=visual_palette_layout_v241(rows,columns=columns,lightness_span=6.0,soft_boundary_deg=4.0,separator_rows=0); new_s=perf_counter()-t2
    if len(new["keys"])!=len(samples): raise RuntimeError(f"完整性检查失败：输入 {len(samples)}，输出 {len(new['keys'])}")
    old_m=_row_metrics(old["grid"],by_key,old["slot_info"]); new_m=_row_metrics(new["grid"],by_key,new["slot_info"])
    old_vert,_=_vertical_nearest_c_metrics(old["grid"],by_key,old["slot_info"]); new_vert,_=_vertical_nearest_c_metrics(new["grid"],by_key,new["slot_info"])
    hue_bridges,neutral_bridges,conflicts,hue_samples,neutral_samples,conflict_samples=_cross_family_near_pairs_v241(new["keys"],by_key,new["slot_info"],3.0)
    low_conf=[k for k in new["keys"] if new["slot_info"][k].get("hue_confidence") in {"LOW","VERY_LOW"}]
    boundary=[k for k in new["keys"] if new["slot_info"][k].get("boundary_zone")]
    neutral_kinds=Counter(new["slot_info"][k].get("neutral_kind","") for k in new["keys"] if new["slot_info"][k]["family"]=="Neutral Field")
    neutral_tints=Counter(new["slot_info"][k].get("neutral_tint","") for k in new["keys"] if new["slot_info"][k]["family"]=="Neutral Field")

    out_dir.mkdir(parents=True,exist_ok=True); stamp=datetime.now().strftime("%Y%m%d_%H%M%S"); safe_stem="".join(ch if ch.isalnum() or ch in "-_" else "_" for ch in qtx.stem)[:80] or "qtx"; base=f"{prefix}_{safe_stem}_{stamp}"
    xlsx_path=out_dir/f"{base}_AUDIT.xlsx"; png_path=out_dir/f"{base}_PREVIEW.png"; txt_path=out_dir/f"{base}_AUDIT.txt"; json_path=out_dir/f"{base}_AUDIT.json"
    _write_preview_png(png_path,new["grid"],by_key,new["slot_info"],columns)

    wb=Workbook(); ws=wb.active; ws.title="总览"
    summary=[
        ("项目","结果 / 定义"),("输入文件",str(qtx)),("输入色样",len(samples)),("输出色样",len(new["keys"])),("完整性","PASS"),
        ("QTX解析时间(s)",round(parse_s,4)),("V2.4时间(s)",round(old_s,4)),("V2.4.1时间(s)",round(new_s,4)),
        ("V2.4 行内 ΔE00 mean / p95 / max",f"{old_m['de']['mean']:.3f} / {old_m['de']['p95']:.3f} / {old_m['de']['max']:.3f}"),
        ("V2.4.1 行内 ΔE00 mean / p95 / max",f"{new_m['de']['mean']:.3f} / {new_m['de']['p95']:.3f} / {new_m['de']['max']:.3f}"),
        ("V2.4.1 行内 |ΔL*| mean / p95 / max",f"{new_m['L']['mean']:.3f} / {new_m['L']['p95']:.3f} / {new_m['L']['max']:.3f}"),
        ("V2.4.1 行内 |ΔC*| mean / p95 / max",f"{new_m['C_step']['mean']:.3f} / {new_m['C_step']['p95']:.3f} / {new_m['C_step']['max']:.3f}"),
        ("V2.4.1 行最大C*跨度 mean / p95 / max",f"{new_m['row_c_span']['mean']:.3f} / {new_m['row_c_span']['p95']:.3f} / {new_m['row_c_span']['max']:.3f}"),
        ("V2.4.1 C*单调违规",new_m["c_monotonic_violations"]),
        ("V2.4 垂直近C邻接 ΔE00 mean / p95 / max",f"{old_vert['mean']:.3f} / {old_vert['p95']:.3f} / {old_vert['max']:.3f}"),
        ("V2.4.1 垂直近C邻接 ΔE00 mean / p95 / max",f"{new_vert['mean']:.3f} / {new_vert['p95']:.3f} / {new_vert['max']:.3f}"),
        ("相邻Hue Family软边界桥接对(ΔE00<3)",len(hue_bridges)),("Hue桥接涉及唯一色样",len(hue_samples)),
        ("Neutral Field过渡桥接对(ΔE00<3)",len(neutral_bridges)),("Neutral桥接涉及唯一色样",len(neutral_samples)),
        ("真正语义冲突对(ΔE00<3)",len(conflicts)),("语义冲突涉及唯一色样",len(conflict_samples)),
        ("低Hue置信度彩色样(C*<12)",len(low_conf)),("Hue边界区色样(距30°分界≤4°)",len(boundary)),
        ("Neutral Field 色样",new["neutral_count"]),("Neutral Field构成",dict(neutral_kinds)),("Neutral Tint分布",dict(neutral_tints)),
        ("Near-White自适应桥接吸收",new.get("near_white_bridge_count",0)),
        ("V2.4.1定位","离线验证；保持V2.4全部规则，仅对同数据集中与Tinted White达到ΔE00<3的高L*近白彩色做自适应桥接吸收；尚未进入正式色卡编排。"),
    ]
    for r,(a,b) in enumerate(summary,1):
        ws.cell(r,1,a); ws.cell(r,2,str(b) if isinstance(b,dict) else b)
        if r==1: ws.cell(r,1).font=ws.cell(r,2).font=Font(bold=True); ws.cell(r,1).fill=ws.cell(r,2).fill=PatternFill("solid",fgColor="E8EEF7")
    ws.column_dimensions["A"].width=48; ws.column_dimensions["B"].width=100; ws.sheet_view.showGridLines=False

    _write_palette_sheet(wb,"V2.4.1二维编排",new["grid"],by_key,new["slot_info"],columns)
    det=wb.create_sheet("V2.4.1数据明细"); headers=["顺序","Grid Row","Grid Col","色块","名称","Family","Neutral Kind","Neutral Tint","Hue Confidence","Boundary Zone","Adjacent Family","Near-White Bridge","Bridge ΔE00","Bridge Seed","From Family","L*","a*","b*","C*","h°","L Shelf","C Row"]
    for c,h in enumerate(headers,1): det.cell(1,c,h).font=Font(bold=True); det.cell(1,c).fill=PatternFill("solid",fgColor="E8EEF7")
    rank=0
    for r_idx,row in enumerate(new["grid"]):
        for c_idx,key in enumerate(row):
            if key is None: continue
            rank+=1; sm=by_key[key]; lab=tuple(float(x) for x in sm.lab_d65_10); L,a,b=lab; C=math.hypot(a,b); h=math.degrees(math.atan2(b,a))%360 if C>1e-12 else 0.0; inf=new["slot_info"][key]
            vals=[rank,r_idx+1,c_idx+1,"",sm.display_name or key,inf["family"],inf.get("neutral_kind",""),inf.get("neutral_tint",""),inf.get("hue_confidence",""),"YES" if inf.get("boundary_zone") else "",inf.get("adjacent_family",""),"YES" if inf.get("near_white_bridge") else "",inf.get("bridge_de00",0.0),inf.get("bridge_seed",""),inf.get("bridge_from_family",""),L,a,b,C,h,inf.get("lightness_band",""),inf.get("chroma_row","")]
            for col,v in enumerate(vals,1): det.cell(rank+1,col,v)
            det.cell(rank+1,4).fill=PatternFill("solid",fgColor=_safe_hex(lab))
    det.freeze_panes="A2"; det.auto_filter.ref=f"A1:V{rank+1}"; [setattr(det.column_dimensions[get_column_letter(col)],'width',14) for col in range(1,23)]; det.column_dimensions["E"].width=45; det.column_dimensions["F"].width=20; det.column_dimensions["G"].width=22; det.column_dimensions["H"].width=20; det.column_dimensions["N"].width=40; det.column_dimensions["O"].width=20

    qm=wb.create_sheet("行质量指标"); qheaders=["Grid Row","数量","Family","L* min","L* max","L* span","C* min","C* max","C* span","行内ΔE00 mean","行内ΔE00 max","行内hue jump max"]
    for c,h in enumerate(qheaders,1): qm.cell(1,c,h).font=Font(bold=True); qm.cell(1,c).fill=PatternFill("solid",fgColor="E8EEF7")
    for rr,row in enumerate(new_m["rows"],2):
        vals=[row["row"],row["count"],row["family"],row["L_min"],row["L_max"],row["L_span"],row["C_min"],row["C_max"],row["C_span"],row["de_mean"],row["de_max"],row["hue_jump_max"]]
        for c,v in enumerate(vals,1): qm.cell(rr,c,v)
    qm.freeze_panes="A2"; qm.auto_filter.ref=f"A1:L{len(new_m['rows'])+1}"

    bridge_keys=[k for k in new["keys"] if new["slot_info"][k].get("near_white_bridge")]
    br=wb.create_sheet("NearWhite桥接吸收"); bheaders=["样本","原Family","新Family","ΔE00 to Seed","Seed样本","L*","C*","h°"]
    for c,h in enumerate(bheaders,1): br.cell(1,c,h).font=Font(bold=True); br.cell(1,c).fill=PatternFill("solid",fgColor="E8F5E9")
    for rr,key in enumerate(bridge_keys,2):
        sm=by_key[key]; lab=tuple(float(x) for x in sm.lab_d65_10); L,a,b=lab; C=math.hypot(a,b); h=math.degrees(math.atan2(b,a))%360 if C>1e-12 else 0.0; inf=new["slot_info"][key]
        vals=[sm.display_name or key,inf.get("bridge_from_family",""),inf.get("family",""),inf.get("bridge_de00",0.0),inf.get("bridge_seed",""),L,C,h]
        for c,v in enumerate(vals,1): br.cell(rr,c,v)
    br.freeze_panes="A2"; br.auto_filter.ref=f"A1:H{max(2,len(bridge_keys)+1)}"; br.column_dimensions["A"].width=45; br.column_dimensions["E"].width=45

    _write_pairs(wb,"Hue软边界桥接",hue_bridges,by_key,"E8F5E9","相邻Hue Family且视觉很接近：保留为软边界桥接")
    _write_pairs(wb,"Neutral过渡桥接",neutral_bridges,by_key,"FFF7D6","Neutral Field 与低Hue置信度彩色样视觉很接近：属于连续过渡，不计为错误")
    _write_pairs(wb,"语义冲突",conflicts,by_key,"FDECEC","视觉很接近但属于非相邻的强色相Family：优先人工复核")
    wb.save(xlsx_path)

    payload={"schema":1,"version":"V2.4.1","input":str(qtx),"sample_count":len(samples),"parse_s":parse_s,
             "v24":{"arrange_s":old_s,"row_de":old_m["de"],"vertical_nearest_c":old_vert},
             "v241":{"arrange_s":new_s,"row_de":new_m["de"],"row_L":new_m["L"],"row_C_step":new_m["C_step"],"row_C_span":new_m["row_c_span"],"vertical_nearest_c":new_vert,"c_monotonic_violations":new_m["c_monotonic_violations"],"hue_bridges":len(hue_bridges),"neutral_bridges":len(neutral_bridges),"semantic_conflicts":len(conflicts),"semantic_conflict_unique_samples":len(conflict_samples),"low_hue_confidence":len(low_conf),"neutral_field":new["neutral_count"],"neutral_kinds":dict(neutral_kinds),"neutral_tints":dict(neutral_tints),"family_counts":new["family_counts"],"near_white_bridge_count":new.get("near_white_bridge_count",0)}}
    json_path.write_text(json.dumps(payload,ensure_ascii=False,indent=2),encoding="utf-8")
    lines=["Chromatic Analysis · 视觉色卡编排 V2.4.1 离线验证","="*72,f"文件: {qtx}",f"输入/输出: {len(samples)} / {len(new['keys'])}  PASS",f"V2.4 行内ΔE00: mean={old_m['de']['mean']:.3f}, p95={old_m['de']['p95']:.3f}, max={old_m['de']['max']:.3f}",f"V2.4.1 行内ΔE00: mean={new_m['de']['mean']:.3f}, p95={new_m['de']['p95']:.3f}, max={new_m['de']['max']:.3f}",f"V2.4.1 垂直近C邻接ΔE00: mean={new_vert['mean']:.3f}, p95={new_vert['p95']:.3f}, max={new_vert['max']:.3f}",f"Near-White桥接吸收: {new.get('near_white_bridge_count',0)}",f"Hue软边界桥接: {len(hue_bridges)}",f"Neutral过渡桥接: {len(neutral_bridges)}",f"真正语义冲突: {len(conflicts)} 对 / {len(conflict_samples)} 个唯一色样",f"Neutral Field: {new['neutral_count']}  {dict(neutral_kinds)}","",f"预览: {png_path.name}",f"Excel: {xlsx_path.name}"]
    txt_path.write_text("\n".join(lines),encoding="utf-8")
    return {"qtx":str(qtx),"sample_count":len(samples),"v24_de":old_m["de"],"v241_de":new_m["de"],"v24_vert":old_vert,"v241_vert":new_vert,"v241_cspan":new_m["row_c_span"],"near_white_bridge_count":new.get("near_white_bridge_count",0),"hue_bridges":len(hue_bridges),"neutral_bridges":len(neutral_bridges),"conflicts":len(conflicts),"conflict_samples":len(conflict_samples),"neutral":new["neutral_count"],"xlsx":str(xlsx_path),"png":str(png_path),"txt":str(txt_path),"json":str(json_path)}


def _batch(folder: Path, columns: int) -> int:
    qtx_files=sorted(p for p in folder.rglob("*.qtx") if p.is_file())
    if not qtx_files: print(f"文件夹中没有找到 QTX: {folder}"); return 2
    stamp=datetime.now().strftime("%Y%m%d_%H%M%S"); batch_dir=ROOT/"diagnostics_reports"/f"VPA_V241_BATCH_{stamp}"; batch_dir.mkdir(parents=True,exist_ok=True); results=[]
    print(f"批量验证 {len(qtx_files)} 个 QTX；结果目录: {batch_dir}")
    for i,qtx in enumerate(qtx_files,1):
        print(f"[{i}/{len(qtx_files)}] {qtx.name}")
        try: results.append(_audit_one(qtx,batch_dir,columns=columns,prefix=f"{i:03d}_VPA_V241"))
        except Exception as exc: print(f"    失败: {exc}"); results.append({"qtx":str(qtx),"error":repr(exc)})
    wb=Workbook(); ws=wb.active; ws.title="批量汇总"; headers=["文件","色样","V2.4 ΔE mean","V2.4.1 ΔE mean","V2.4 ΔE p95","V2.4.1 ΔE p95","V2.4 ΔE max","V2.4.1 ΔE max","V2.4 垂直近C mean","V2.4.1 垂直近C mean","V2.4.1 行C跨度 p95","Near-White桥接吸收","Hue桥接","Neutral桥接","真正语义冲突","冲突唯一色样","Neutral Field","状态"]
    for c,h in enumerate(headers,1): ws.cell(1,c,h).font=Font(bold=True); ws.cell(1,c).fill=PatternFill("solid",fgColor="E8EEF7")
    for r,item in enumerate(results,2):
        if "error" in item: vals=[item["qtx"]]+[""]*16+[item["error"]]
        else: vals=[item["qtx"],item["sample_count"],item["v24_de"]["mean"],item["v241_de"]["mean"],item["v24_de"]["p95"],item["v241_de"]["p95"],item["v24_de"]["max"],item["v241_de"]["max"],item["v24_vert"]["mean"],item["v241_vert"]["mean"],item["v241_cspan"]["p95"],item["near_white_bridge_count"],item["hue_bridges"],item["neutral_bridges"],item["conflicts"],item["conflict_samples"],item["neutral"],"PASS"]
        for c,v in enumerate(vals,1): ws.cell(r,c,v)
    ws.freeze_panes="A2"; ws.auto_filter.ref=f"A1:R{len(results)+1}"; ws.column_dimensions["A"].width=80
    summary_xlsx=ROOT/"diagnostics_reports"/f"VPA_V241_BATCH_SUMMARY_{stamp}.xlsx"; wb.save(summary_xlsx)
    summary_json=ROOT/"diagnostics_reports"/f"VPA_V241_BATCH_SUMMARY_{stamp}.json"; summary_json.write_text(json.dumps(results,ensure_ascii=False,indent=2),encoding="utf-8")
    print("\n批量验证完成。")
    print(f"汇总 Excel: {summary_xlsx}\n汇总 JSON: {summary_json}")
    print("建议先把批量汇总 Excel 发给 ChatGPT；再按需要补代表性 PNG + XLSX。")
    return 0


def main() -> int:
    ap=argparse.ArgumentParser(description="Chromatic Analysis Visual Palette V2.4.1 offline audit"); ap.add_argument("path",type=Path,help="QTX file or a folder containing QTX files"); ap.add_argument("--columns",type=int,default=6); args=ap.parse_args(); path=args.path.expanduser().resolve()
    if not path.exists(): print(f"路径不存在: {path}"); return 2
    print("Chromatic Analysis · 视觉色卡编排 V2.4.1 离线验证"); print("="*72)
    print("V2.4.1: V2.4 + Data-Adaptive Near-White Bridge Absorption")
    print("说明：只生成报告/预览，不修改正式色卡编排。\n")
    if path.is_dir(): return _batch(path,args.columns)
    if path.suffix.lower()!=".qtx": print("请选择 .qtx 文件，或包含 .qtx 的文件夹。"); return 2
    result=_audit_one(path,ROOT/"diagnostics_reports",columns=args.columns)
    print("完成。"); print(f"预览: {result['png']}\nExcel: {result['xlsx']}\nTXT: {result['txt']}\nJSON: {result['json']}"); print("请把 PNG + XLSX 发给 ChatGPT。"); return 0


if __name__ == "__main__":
    raise SystemExit(main())
