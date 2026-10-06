from __future__ import annotations

import argparse
import json
import math
import statistics
import sys
from datetime import datetime
from pathlib import Path
from time import perf_counter

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

from qtx_core import parse_qtx_file, visual_palette_layout_v21, visual_palette_layout_v22
from qtx_core.colorimetry import delta_e
from qtx_app.excel_exchange import _lab_to_hex


def _pct(vals, q: float) -> float:
    vals = sorted(float(v) for v in vals)
    if not vals:
        return 0.0
    idx = min(len(vals) - 1, max(0, int(round((len(vals) - 1) * q))))
    return vals[idx]


def _summary(vals) -> dict[str, float]:
    vals = [float(v) for v in vals]
    if not vals:
        return {"count": 0, "mean": 0.0, "median": 0.0, "p95": 0.0, "max": 0.0}
    return {
        "count": len(vals),
        "mean": statistics.fmean(vals),
        "median": statistics.median(vals),
        "p95": _pct(vals, 0.95),
        "max": max(vals),
    }


def _thin_border():
    side = Side(style="thin", color="D9E0EA")
    return Border(left=side, right=side, top=side, bottom=side)


def _safe_hex(lab) -> str:
    try:
        return _lab_to_hex(tuple(float(x) for x in lab)).replace("#", "")
    except Exception:
        return "D9D9D9"


def _short_hue_distance(a: float, b: float) -> float:
    d = abs(float(a) - float(b)) % 360.0
    return min(d, 360.0 - d)


def _family_adjacent(i: int, j: int) -> bool:
    if i < 0 or j < 0 or i >= 12 or j >= 12:
        return False
    d = abs(i - j)
    return d == 1 or d == 11


def _row_metrics(grid, by_key, info):
    rows = []
    within_de, within_hue, within_L, within_C = [], [], [], []
    c_monotonic_violations = 0
    for r_idx, row in enumerate(grid):
        keys = [k for k in row if k is not None]
        if not keys:
            continue
        labs = [by_key[k].lab_d65_10 for k in keys]
        Ls = [float(x[0]) for x in labs]
        Cs = [math.hypot(float(x[1]), float(x[2])) for x in labs]
        hs = [math.degrees(math.atan2(float(x[2]), float(x[1]))) % 360.0 if c > 1e-12 else 0.0 for x, c in zip(labs, Cs)]
        des, hds = [], []
        for i in range(len(keys) - 1):
            de = float(delta_e(labs[i], labs[i + 1], "CIE 2000"))
            des.append(de); within_de.append(de)
            within_L.append(abs(Ls[i] - Ls[i + 1]))
            within_C.append(abs(Cs[i] - Cs[i + 1]))
            fam = info[keys[i]]["family"]
            fam2 = info[keys[i+1]]["family"]
            if fam == fam2 and fam != "Neutral Axis" and Cs[i] >= 10.0 and Cs[i+1] >= 10.0:
                hd = _short_hue_distance(hs[i], hs[i + 1])
                hds.append(hd); within_hue.append(hd)
            if fam == fam2 and fam != "Neutral Axis" and Cs[i+1] + 1e-9 < Cs[i]:
                c_monotonic_violations += 1
        fams = [info[k]["family"] for k in keys]
        rows.append({
            "row": r_idx + 1,
            "count": len(keys),
            "family": fams[0] if len(set(fams)) == 1 else " / ".join(dict.fromkeys(fams)),
            "L_min": min(Ls), "L_max": max(Ls), "L_span": max(Ls) - min(Ls),
            "C_min": min(Cs), "C_max": max(Cs), "C_span": max(Cs) - min(Cs),
            "de_mean": statistics.fmean(des) if des else 0.0,
            "de_max": max(des) if des else 0.0,
            "hue_jump_max": max(hds) if hds else 0.0,
        })
    return {
        "rows": rows,
        "de": _summary(within_de),
        "hue": _summary(within_hue),
        "L": _summary(within_L),
        "C_step": _summary(within_C),
        "c_monotonic_violations": c_monotonic_violations,
        "row_c_span": _summary([r["C_span"] for r in rows]),
    }


def _section_transition_metrics(grid, by_key, info):
    nonempty = [(idx, [k for k in row if k is not None]) for idx, row in enumerate(grid) if any(k is not None for k in row)]
    vals, detail = [], []
    for (r1, a), (r2, b) in zip(nonempty, nonempty[1:]):
        if not a or not b:
            continue
        if info[a[0]]["family"] != info[b[0]]["family"]:
            continue
        k1, k2 = a[-1], b[0]
        de = float(delta_e(by_key[k1].lab_d65_10, by_key[k2].lab_d65_10, "CIE 2000"))
        vals.append(de)
        detail.append((r1 + 1, r2 + 1, info[k1]["family"], k1, k2, de))
    return _summary(vals), detail


def _cross_family_near_pairs(keys, by_key, info, threshold=3.0):
    bridges, conflicts = [], []
    involved_bridge, involved_conflict = set(), set()
    for i, ka in enumerate(keys):
        ia = int(info[ka].get("family_index", 99))
        fa = info[ka]["family"]
        for kb in keys[i+1:]:
            fb = info[kb]["family"]
            if fa == fb:
                continue
            de = float(delta_e(by_key[ka].lab_d65_10, by_key[kb].lab_d65_10, "CIE 2000"))
            if de >= threshold:
                continue
            ib = int(info[kb].get("family_index", 99))
            item = (de, ka, kb, fa, fb, ia, ib)
            if _family_adjacent(ia, ib):
                bridges.append(item); involved_bridge.update((ka, kb))
            else:
                conflicts.append(item); involved_conflict.update((ka, kb))
    bridges.sort(key=lambda x: x[0]); conflicts.sort(key=lambda x: x[0])
    return bridges, conflicts, involved_bridge, involved_conflict


def _write_preview_png(path: Path, grid, by_key, info, columns: int) -> bool:
    try:
        from PIL import Image, ImageDraw, ImageFont
    except Exception:
        return False
    cell_w, cell_h, sw_h = 220, 112, 66
    left_label = 120
    width = left_label + columns * cell_w
    height = max(1, len(grid)) * cell_h
    im = Image.new("RGB", (width, height), "white")
    draw = ImageDraw.Draw(im)
    font = small = None
    for fp in (Path("C:/Windows/Fonts/msyh.ttc"), Path("C:/Windows/Fonts/msyh.ttf"), Path("C:/Windows/Fonts/simhei.ttf")):
        if fp.exists():
            try:
                font = ImageFont.truetype(str(fp), 14); small = ImageFont.truetype(str(fp), 11); break
            except Exception:
                pass
    if font is None:
        font = ImageFont.load_default(); small = font
    last_family = None
    rank = 0
    for r_idx, row in enumerate(grid):
        y = r_idx * cell_h
        keys = [k for k in row if k is not None]
        if not keys:
            continue
        family = info[keys[0]]["family"]
        if family != last_family:
            draw.line((0, y, width, y), fill=(180, 190, 204), width=2)
            draw.text((6, y + 8), family, fill=(45, 55, 72), font=small)
            last_family = family
        for c_idx, key in enumerate(row):
            if key is None:
                continue
            rank += 1
            sm = by_key[key]
            lab = tuple(float(x) for x in sm.lab_d65_10)
            L, a, b = lab; C = math.hypot(a, b)
            h = math.degrees(math.atan2(b, a)) % 360.0 if C > 1e-12 else 0.0
            hx = _safe_hex(lab); rgb = tuple(int(hx[i:i+2], 16) for i in (0, 2, 4))
            x = left_label + c_idx * cell_w
            draw.rounded_rectangle((x + 4, y + 4, x + cell_w - 8, y + sw_h), radius=8, fill=rgb, outline=(215, 220, 228))
            name = sm.display_name or "(unnamed)"
            if len(name) > 24: name = name[:23] + "…"
            mark = " ◇" if info[key].get("boundary_zone") else ""
            draw.text((x + 7, y + sw_h + 5), f"{rank:03d} {name}{mark}", fill=(20, 25, 32), font=small)
            draw.text((x + 7, y + sw_h + 24), f"L {L:.1f}  C {C:.1f}  h {h:.0f}°", fill=(85, 95, 110), font=small)
    im.save(path)
    return True


def _write_palette_sheet(wb, title, grid, by_key, info, columns):
    ws = wb.create_sheet(title); ws.sheet_view.showGridLines = False
    border = _thin_border(); last_family = None; out_row = 1
    for row in grid:
        keys = [k for k in row if k is not None]
        if not keys: continue
        family = info[keys[0]]["family"]
        if family != last_family:
            ws.cell(out_row, 1, family); ws.cell(out_row, 1).font = Font(bold=True, color="334155")
            ws.merge_cells(start_row=out_row, start_column=1, end_row=out_row, end_column=columns)
            out_row += 1; last_family = family
        for c_idx, key in enumerate(row, 1):
            if key is None: continue
            sm = by_key[key]; lab = tuple(float(x) for x in sm.lab_d65_10)
            label = sm.display_name or key
            if info[key].get("boundary_zone"): label += "  ◇边界"
            cell = ws.cell(out_row, c_idx, label)
            cell.fill = PatternFill("solid", fgColor=_safe_hex(lab)); cell.border = border
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True); cell.font = Font(size=9)
        ws.row_dimensions[out_row].height = 48; out_row += 1
    for c in range(1, columns + 1): ws.column_dimensions[get_column_letter(c)].width = 24
    return ws


def _audit_one(qtx: Path, out_dir: Path, *, columns: int = 6, prefix: str = "VPA_V22") -> dict:
    t0 = perf_counter(); samples = parse_qtx_file(qtx); parse_s = perf_counter() - t0
    rows, by_key = [], {}
    for i, sm in enumerate(samples, 1):
        key = f"S{i:04d}"; by_key[key] = sm; rows.append((key, sm.display_name or key, sm.lab_d65_10))

    t1 = perf_counter(); old = visual_palette_layout_v21(rows, columns=columns, lightness_span=7.5, separator_rows=0); old_s = perf_counter() - t1
    t2 = perf_counter(); new = visual_palette_layout_v22(rows, columns=columns, lightness_span=7.5, chroma_span=25.0, soft_boundary_deg=4.0, separator_rows=0); new_s = perf_counter() - t2
    if len(new["keys"]) != len(samples):
        raise RuntimeError(f"完整性检查失败：输入 {len(samples)}，输出 {len(new['keys'])}")

    old_m = _row_metrics(old["grid"], by_key, old["slot_info"])
    new_m = _row_metrics(new["grid"], by_key, new["slot_info"])
    old_trans, _ = _section_transition_metrics(old["grid"], by_key, old["slot_info"])
    new_trans, _ = _section_transition_metrics(new["grid"], by_key, new["slot_info"])
    bridges, conflicts, bridge_samples, conflict_samples = _cross_family_near_pairs(new["keys"], by_key, new["slot_info"], 3.0)
    low_conf = [k for k in new["keys"] if new["slot_info"][k].get("hue_confidence") in {"LOW", "VERY_LOW"}]
    boundary = [k for k in new["keys"] if new["slot_info"][k].get("boundary_zone")]

    out_dir.mkdir(parents=True, exist_ok=True)
    stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    safe_stem = "".join(ch if ch.isalnum() or ch in "-_" else "_" for ch in qtx.stem)[:80] or "qtx"
    base = f"{prefix}_{safe_stem}_{stamp}"
    xlsx_path = out_dir / f"{base}_AUDIT.xlsx"
    png_path = out_dir / f"{base}_PREVIEW.png"
    txt_path = out_dir / f"{base}_AUDIT.txt"
    json_path = out_dir / f"{base}_AUDIT.json"

    _write_preview_png(png_path, new["grid"], by_key, new["slot_info"], columns)

    wb = Workbook(); ws = wb.active; ws.title = "总览"
    summary = [
        ("项目", "结果 / 定义"), ("输入文件", str(qtx)), ("输入色样", len(samples)), ("输出色样", len(new["keys"])), ("完整性", "PASS"),
        ("QTX解析时间(s)", round(parse_s, 4)), ("V2.1时间(s)", round(old_s, 4)), ("V2.2时间(s)", round(new_s, 4)),
        ("V2.1 行内 ΔE00 mean / p95 / max", f"{old_m['de']['mean']:.3f} / {old_m['de']['p95']:.3f} / {old_m['de']['max']:.3f}"),
        ("V2.2 行内 ΔE00 mean / p95 / max", f"{new_m['de']['mean']:.3f} / {new_m['de']['p95']:.3f} / {new_m['de']['max']:.3f}"),
        ("V2.2 行内 |ΔL*| mean / p95 / max", f"{new_m['L']['mean']:.3f} / {new_m['L']['p95']:.3f} / {new_m['L']['max']:.3f}"),
        ("V2.2 行内 |ΔC*| mean / p95 / max", f"{new_m['C_step']['mean']:.3f} / {new_m['C_step']['p95']:.3f} / {new_m['C_step']['max']:.3f}"),
        ("V2.2 行最大C*跨度 mean / p95 / max", f"{new_m['row_c_span']['mean']:.3f} / {new_m['row_c_span']['p95']:.3f} / {new_m['row_c_span']['max']:.3f}"),
        ("V2.2 C*单调违规", new_m["c_monotonic_violations"]),
        ("V2.1 同Family行间 ΔE00 mean / p95 / max", f"{old_trans['mean']:.3f} / {old_trans['p95']:.3f} / {old_trans['max']:.3f}"),
        ("V2.2 同Family行间 ΔE00 mean / p95 / max", f"{new_trans['mean']:.3f} / {new_trans['p95']:.3f} / {new_trans['max']:.3f}"),
        ("相邻Hue Family软边界桥接对(ΔE00<3)", len(bridges)),
        ("软边界桥接涉及唯一色样", len(bridge_samples)),
        ("非相邻/中性轴语义冲突对(ΔE00<3)", len(conflicts)),
        ("语义冲突涉及唯一色样", len(conflict_samples)),
        ("低Hue置信度彩色样(C*<10)", len(low_conf)),
        ("Hue边界区色样(距30°分界≤4°)", len(boundary)),
        ("Neutral Axis 色样", new["neutral_count"]),
        ("V2.2定位", "离线验证；12色相骨架 + 自适应L*层 + C*横轴 + Hue软边界 + 连续中性轴；尚未进入正式色卡编排。"),
    ]
    for r, (a, b) in enumerate(summary, 1):
        ws.cell(r, 1, a); ws.cell(r, 2, b)
        if r == 1:
            ws.cell(r, 1).font = ws.cell(r, 2).font = Font(bold=True)
            ws.cell(r, 1).fill = ws.cell(r, 2).fill = PatternFill("solid", fgColor="E8EEF7")
    ws.column_dimensions["A"].width = 42; ws.column_dimensions["B"].width = 95; ws.sheet_view.showGridLines = False

    _write_palette_sheet(wb, "V2.2二维编排", new["grid"], by_key, new["slot_info"], columns)

    det = wb.create_sheet("V2.2数据明细")
    headers = ["顺序", "Grid Row", "Grid Col", "色块", "名称", "Family", "Neutral Kind", "Hue Confidence", "Boundary Zone", "Adjacent Family", "L*", "a*", "b*", "C*", "h°", "L Shelf", "C Row"]
    for c, h in enumerate(headers, 1): det.cell(1, c, h).font = Font(bold=True); det.cell(1, c).fill = PatternFill("solid", fgColor="E8EEF7")
    rank = 0
    for r_idx, row in enumerate(new["grid"]):
        for c_idx, key in enumerate(row):
            if key is None: continue
            rank += 1; sm = by_key[key]; lab = tuple(float(x) for x in sm.lab_d65_10); L, a, b = lab; C = math.hypot(a, b); h = math.degrees(math.atan2(b, a)) % 360 if C > 1e-12 else 0.0; inf = new["slot_info"][key]
            vals = [rank, r_idx + 1, c_idx + 1, "", sm.display_name or key, inf["family"], inf.get("neutral_kind", ""), inf.get("hue_confidence", ""), "YES" if inf.get("boundary_zone") else "", inf.get("adjacent_family", ""), L, a, b, C, h, inf.get("lightness_band", ""), inf.get("chroma_row", "")]
            for col, v in enumerate(vals, 1): det.cell(rank + 1, col, v)
            det.cell(rank + 1, 4).fill = PatternFill("solid", fgColor=_safe_hex(lab))
    det.freeze_panes = "A2"; det.auto_filter.ref = f"A1:Q{rank+1}"
    for col in range(1, 18): det.column_dimensions[get_column_letter(col)].width = 14
    det.column_dimensions["E"].width = 45; det.column_dimensions["F"].width = 20; det.column_dimensions["H"].width = 18; det.column_dimensions["J"].width = 20

    qm = wb.create_sheet("行质量指标")
    qheaders = ["Grid Row", "数量", "Family", "L* min", "L* max", "L* span", "C* min", "C* max", "C* span", "行内ΔE00 mean", "行内ΔE00 max", "行内hue jump max"]
    for c, h in enumerate(qheaders, 1): qm.cell(1, c, h).font = Font(bold=True); qm.cell(1, c).fill = PatternFill("solid", fgColor="E8EEF7")
    for rr, row in enumerate(new_m["rows"], 2):
        vals = [row["row"], row["count"], row["family"], row["L_min"], row["L_max"], row["L_span"], row["C_min"], row["C_max"], row["C_span"], row["de_mean"], row["de_max"], row["hue_jump_max"]]
        for c, v in enumerate(vals, 1): qm.cell(rr, c, v)
    qm.freeze_panes = "A2"; qm.auto_filter.ref = f"A1:L{len(new_m['rows'])+1}"

    def write_pairs(title, pairs, fill):
        sh = wb.create_sheet(title)
        headers = ["ΔE00", "样本A", "Family A", "样本B", "Family B", "说明"]
        for c, h in enumerate(headers, 1): sh.cell(1, c, h).font = Font(bold=True); sh.cell(1, c).fill = PatternFill("solid", fgColor=fill)
        for r, (de, ka, kb, fa, fb, ia, ib) in enumerate(pairs, 2):
            sa, sb = by_key[ka], by_key[kb]
            if title == "软边界桥接": note = "相邻Hue Family且视觉很接近：这是软边界桥接，不计为错误"
            else: note = "视觉很接近但属于非相邻Family或Neutral Axis，优先人工复核"
            vals = [de, sa.display_name or ka, fa, sb.display_name or kb, fb, note]
            for c, v in enumerate(vals, 1): sh.cell(r, c, v)
        sh.freeze_panes = "A2"; sh.auto_filter.ref = f"A1:F{max(2,len(pairs)+1)}"; sh.column_dimensions["B"].width = 45; sh.column_dimensions["D"].width = 45; sh.column_dimensions["F"].width = 58
    write_pairs("软边界桥接", bridges, "E8F5E9")
    write_pairs("语义冲突", conflicts, "FDECEC")

    wb.save(xlsx_path)

    payload = {
        "schema": 1, "version": "V2.2", "input": str(qtx), "sample_count": len(samples), "parse_s": parse_s,
        "v21": {"arrange_s": old_s, "row_de": old_m["de"], "row_L": old_m["L"], "same_family_row_transition": old_trans},
        "v22": {"arrange_s": new_s, "row_de": new_m["de"], "row_L": new_m["L"], "row_C_step": new_m["C_step"], "row_C_span": new_m["row_c_span"], "same_family_row_transition": new_trans, "c_monotonic_violations": new_m["c_monotonic_violations"], "soft_boundary_bridges": len(bridges), "soft_boundary_unique_samples": len(bridge_samples), "semantic_conflicts": len(conflicts), "semantic_conflict_unique_samples": len(conflict_samples), "low_hue_confidence": len(low_conf), "boundary_zone_samples": len(boundary), "neutral_axis": new["neutral_count"], "family_counts": new["family_counts"]},
    }
    json_path.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    lines = [
        "Chromatic Analysis · 视觉色卡编排 V2.2 离线验证", "=" * 72, f"文件: {qtx}", f"输入/输出: {len(samples)} / {len(new['keys'])}  PASS",
        f"V2.1 行内ΔE00: mean={old_m['de']['mean']:.3f}, p95={old_m['de']['p95']:.3f}, max={old_m['de']['max']:.3f}",
        f"V2.2 行内ΔE00: mean={new_m['de']['mean']:.3f}, p95={new_m['de']['p95']:.3f}, max={new_m['de']['max']:.3f}",
        f"V2.2 行内|ΔC*|: mean={new_m['C_step']['mean']:.3f}, p95={new_m['C_step']['p95']:.3f}, max={new_m['C_step']['max']:.3f}",
        f"V2.2 行C*跨度: mean={new_m['row_c_span']['mean']:.3f}, p95={new_m['row_c_span']['p95']:.3f}, max={new_m['row_c_span']['max']:.3f}",
        f"V2.2 C*单调违规: {new_m['c_monotonic_violations']}",
        f"软边界桥接: {len(bridges)} 对 / {len(bridge_samples)} 个唯一色样", f"语义冲突: {len(conflicts)} 对 / {len(conflict_samples)} 个唯一色样",
        f"Neutral Axis: {new['neutral_count']}", "", f"预览: {png_path.name}", f"Excel: {xlsx_path.name}"
    ]
    txt_path.write_text("\n".join(lines), encoding="utf-8")
    return {"qtx": str(qtx), "sample_count": len(samples), "parse_s": parse_s, "v21_s": old_s, "v22_s": new_s, "v21_de": old_m["de"], "v22_de": new_m["de"], "v22_cspan": new_m["row_c_span"], "v22_cstep": new_m["C_step"], "v21_trans": old_trans, "v22_trans": new_trans, "bridges": len(bridges), "bridge_samples": len(bridge_samples), "conflicts": len(conflicts), "conflict_samples": len(conflict_samples), "low_conf": len(low_conf), "neutral": new["neutral_count"], "xlsx": str(xlsx_path), "png": str(png_path), "txt": str(txt_path), "json": str(json_path)}


def _batch(folder: Path, columns: int) -> int:
    qtx_files = sorted(p for p in folder.rglob("*.qtx") if p.is_file())
    if not qtx_files:
        print(f"文件夹中没有找到 QTX: {folder}"); return 2
    stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    batch_dir = ROOT / "diagnostics_reports" / f"VPA_V22_BATCH_{stamp}"
    batch_dir.mkdir(parents=True, exist_ok=True)
    results = []
    print(f"批量验证 {len(qtx_files)} 个 QTX；结果目录: {batch_dir}")
    for i, qtx in enumerate(qtx_files, 1):
        print(f"[{i}/{len(qtx_files)}] {qtx.name}")
        try:
            results.append(_audit_one(qtx, batch_dir, columns=columns, prefix=f"{i:03d}_VPA_V22"))
        except Exception as exc:
            print(f"    失败: {exc}")
            results.append({"qtx": str(qtx), "error": repr(exc)})

    wb = Workbook(); ws = wb.active; ws.title = "批量汇总"
    headers = ["文件", "色样", "V2.1 ΔE mean", "V2.2 ΔE mean", "V2.1 ΔE p95", "V2.2 ΔE p95", "V2.1 ΔE max", "V2.2 ΔE max", "V2.2 行C跨度 p95", "软边界桥接对", "语义冲突对", "语义冲突唯一色样", "Neutral Axis", "状态"]
    for c, h in enumerate(headers, 1): ws.cell(1, c, h).font = Font(bold=True); ws.cell(1, c).fill = PatternFill("solid", fgColor="E8EEF7")
    for r, item in enumerate(results, 2):
        if "error" in item:
            vals = [item["qtx"]] + [""] * 12 + [item["error"]]
        else:
            vals = [item["qtx"], item["sample_count"], item["v21_de"]["mean"], item["v22_de"]["mean"], item["v21_de"]["p95"], item["v22_de"]["p95"], item["v21_de"]["max"], item["v22_de"]["max"], item["v22_cspan"]["p95"], item["bridges"], item["conflicts"], item["conflict_samples"], item["neutral"], "PASS"]
        for c, v in enumerate(vals, 1): ws.cell(r, c, v)
    ws.freeze_panes = "A2"; ws.auto_filter.ref = f"A1:N{len(results)+1}"; ws.column_dimensions["A"].width = 80
    summary_xlsx = ROOT / "diagnostics_reports" / f"VPA_V22_BATCH_SUMMARY_{stamp}.xlsx"
    wb.save(summary_xlsx)
    summary_json = ROOT / "diagnostics_reports" / f"VPA_V22_BATCH_SUMMARY_{stamp}.json"
    summary_json.write_text(json.dumps(results, ensure_ascii=False, indent=2), encoding="utf-8")
    print("\n批量验证完成。")
    print(f"汇总 Excel: {summary_xlsx}\n汇总 JSON: {summary_json}")
    print("建议先把汇总 Excel 发给 ChatGPT；再挑表现异常/代表性数据的 PNG + XLSX 继续分析。")
    return 0


def main() -> int:
    ap = argparse.ArgumentParser(description="Chromatic Analysis Visual Palette V2.2 offline audit")
    ap.add_argument("path", type=Path, help="QTX file or a folder containing QTX files")
    ap.add_argument("--columns", type=int, default=6)
    args = ap.parse_args(); path = args.path.expanduser().resolve()
    if not path.exists():
        print(f"路径不存在: {path}"); return 2

    print("Chromatic Analysis · 视觉色卡编排 V2.2 离线验证")
    print("=" * 72)
    print("V2.2: Hue Family + Adaptive Lightness + Chroma Axis + Soft Boundary + Neutral Continuum")
    print("说明：只生成报告/预览，不修改正式色卡编排。")
    print()

    if path.is_dir():
        return _batch(path, args.columns)

    if path.suffix.lower() != ".qtx":
        print("请选择 .qtx 文件，或包含 .qtx 的文件夹。"); return 2
    result = _audit_one(path, ROOT / "diagnostics_reports", columns=args.columns)
    print("完成。")
    print(f"预览: {result['png']}\nExcel: {result['xlsx']}\nTXT: {result['txt']}\nJSON: {result['json']}")
    print("请把 PNG + XLSX 发给 ChatGPT。")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
