from __future__ import annotations

import math
import uuid
from dataclasses import dataclass, field
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

from qtx_core.models import Sample


def short_code(name: str) -> str:
    parts = str(name or "").strip().split()
    return parts[0].upper() if parts else ""


@dataclass
class ExcelImportResult:
    samples: list[Sample] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)
    layout: list[dict] = field(default_factory=list)
    internal_fields_present: bool = False


def _rows_by_code(sheet, name_column: int = 2) -> dict[str, tuple]:
    result = {}
    for row in sheet.iter_rows(min_row=5, values_only=True):
        code = short_code(row[name_column - 1] if len(row) >= name_column else "")
        # MPC files often use 4P codes, but customer workbooks may use other
        # stable identifiers (for example TEST-1 or P025-A1).
        if code:
            result[code] = row
    return result


def import_mpc_data_workbook(path: str) -> ExcelImportResult:
    """读取MPC三表数据Excel：坐标、反射率、内部属性。"""
    wb = load_workbook(path, data_only=True, read_only=True)
    required = {"CIE Coordinates", "Reflectance Data", "Attributes"}
    missing = required.difference(wb.sheetnames)
    if missing:
        raise ValueError("缺少工作表：" + "、".join(sorted(missing)))

    cie = _rows_by_code(wb["CIE Coordinates"])
    refl_sheet = wb["Reflectance Data"]
    refl = _rows_by_code(refl_sheet)
    attrs = _rows_by_code(wb["Attributes"])
    header = [cell.value for cell in refl_sheet[4]]
    wavelength_cols = []
    for index, value in enumerate(header):
        text = str(value or "").lower().replace("nm", "").strip()
        if text.isdigit() and 360 <= int(text) <= 700:
            wavelength_cols.append((index, int(text)))
    if len(wavelength_cols) != 35:
        raise ValueError(f"反射率表应包含360–700 nm共35点，实际识别到{len(wavelength_cols)}点")

    output = ExcelImportResult(internal_fields_present=bool(attrs))
    occupied_ranks: dict[tuple[int, int], str] = {}
    all_codes = list(dict.fromkeys([*cie.keys(), *refl.keys(), *attrs.keys()]))
    for code in all_codes:
        if code not in refl:
            output.warnings.append(f"{code}：缺少反射率，未导入为光谱色样")
            continue
        row = refl[code]
        values = []
        try:
            for index, _ in wavelength_cols:
                value = row[index]
                if value is None:
                    raise ValueError("存在空白")
                values.append(float(value))
        except (ValueError, TypeError, IndexError) as exc:
            output.warnings.append(f"{code}：反射率无效（{exc}）")
            continue

        cie_row = cie.get(code)
        if cie_row and len(cie_row) >= 10:
            xyz = tuple(float(v) for v in cie_row[2:5])
            lab = tuple(float(v) for v in cie_row[5:8])
            display_name = str(cie_row[1]).strip()
        else:
            from qtx_core.colorimetry import reflectance_to_xyz_lab
            xyz, lab = reflectance_to_xyz_lab(values, "D65", [x[1] for x in wavelength_cols], 10)
            display_name = str(row[1]).strip()
            output.warnings.append(f"{code}：坐标表无对应记录，已从光谱重新计算")

        raw = {"EXCEL_SOURCE": Path(path).name, "EXCEL_CODE": code}
        attr_row = attrs.get(code)
        if attr_row:
            names = [str(c.value or "") for c in wb["Attributes"][4]]
            for key, value in zip(names, attr_row):
                if key and value not in (None, ""):
                    raw[f"ATTR_{key}"] = str(value)
        sample_id = str(uuid.uuid5(uuid.NAMESPACE_URL, f"{Path(path).resolve()}|{code}"))
        output.samples.append(Sample(
            sample_id=sample_id, display_name=display_name, kind="STD",
            xyz_d65_10=xyz, lab_d65_10=lab, reflectance=tuple(values),
            wavelengths=tuple(x[1] for x in wavelength_cols), source_file=str(Path(path).resolve()),
            viewing="Excel D65 / 10°", raw=raw,
        ))
        rank = raw.get("ATTR_Rank")
        if rank:
            try:
                column, row_number = (int(x) for x in rank.split("-", 1))
                position = (row_number - 1, column - 1)
                if position in occupied_ranks:
                    output.warnings.append(
                        f"Rank {rank}位置冲突：{occupied_ranks[position]} 与 {code}；两条均保留待人工调整")
                else:
                    occupied_ranks[position] = code
                output.layout.append({"sample_id": sample_id, "code": code,
                                      "row": row_number - 1, "column": column - 1})
            except Exception:
                output.warnings.append(f"{code}：Rank“{rank}”无法识别")
    return output



def _as_float(value):
    if value in (None, ""):
        return None
    try:
        number=float(value)
        return number if math.isfinite(number) else None
    except (TypeError,ValueError):
        return None


def _lab_to_xyz_d65_10(lab: tuple[float,float,float]) -> tuple[float,float,float]:
    """Inverse CIELAB used only when an imported workbook has Lab but no XYZ/spectrum."""
    from qtx_core.colorimetry import reference_white_xyz
    L,a,b=(float(lab[0]),float(lab[1]),float(lab[2]))
    Xn,Yn,Zn=reference_white_xyz('D65',10)
    fy=(L+16.0)/116.0; fx=fy+a/500.0; fz=fy-b/200.0
    delta=6.0/29.0
    def inv(t):
        return t**3 if t>delta else 3.0*delta*delta*(t-4.0/29.0)
    return (Xn*inv(fx),Yn*inv(fy),Zn*inv(fz))


def _sheet_header_map(ws, row_number: int) -> dict[str,int]:
    values=next(ws.iter_rows(min_row=row_number,max_row=row_number,values_only=True),())
    return {str(v or '').strip():i for i,v in enumerate(values) if str(v or '').strip()}


def _exported_spectrum_index(wb) -> dict[str,tuple[tuple[int,...],tuple[float,...],str]]:
    """Read the app's 光谱数据 sheet without requiring it to exist."""
    if '光谱数据' not in wb.sheetnames:
        return {}
    ws=wb['光谱数据']; header=next(ws.iter_rows(min_row=1,max_row=1,values_only=True),())
    wave_cols=[]
    for i,value in enumerate(header):
        m=str(value or '').strip().lower().replace('nm','').strip()
        if m.isdigit() and 300<=int(m)<=800:
            wave_cols.append((i,int(m)))
    out={}
    for row in ws.iter_rows(min_row=2,values_only=True):
        name=str(row[0] or '').strip() if row else ''
        if not name or not wave_cols: continue
        waves=[]; values=[]
        for idx,wave in wave_cols:
            value=_as_float(row[idx] if idx<len(row) else None)
            if value is not None:
                waves.append(wave);values.append(value)
        if len(values)>=2:
            source=str(row[1] or '').strip() if len(row)>1 else ''
            out.setdefault(name,(tuple(waves),tuple(values),source))
    return out


def _sample_from_export_row(path: str, sheet_name: str, row_number: int, name: str,
                            lab, xyz=None, kind='STD', spectrum=None, raw=None) -> Sample:
    lab=tuple(float(x) for x in lab)
    spec=tuple(); waves=tuple()
    if spectrum:
        waves,spec,_source=spectrum
    if xyz is None:
        if spec:
            from qtx_core.colorimetry import reflectance_to_xyz_lab
            xyz, lab_from_spec = reflectance_to_xyz_lab(spec,'D65',waves,10)
            # Keep the exported Lab as the user's record when present. Spectrum is
            # preserved independently and can be re-evaluated under other conditions.
        else:
            xyz=_lab_to_xyz_d65_10(lab)
    sample_id=str(uuid.uuid5(uuid.NAMESPACE_URL,f'{Path(path).resolve()}|{sheet_name}|{row_number}|{name}|{kind}'))
    payload={'EXCEL_SOURCE':Path(path).name,'EXCEL_SHEET':sheet_name,'EXCEL_ROW':str(row_number)}
    payload.update({str(k):str(v) for k,v in (raw or {}).items() if v not in (None,'')})
    return Sample(sample_id=sample_id,display_name=name,kind=kind,
                  xyz_d65_10=tuple(float(x) for x in xyz),lab_d65_10=lab,
                  reflectance=tuple(float(x) for x in spec),wavelengths=tuple(int(x) for x in waves),
                  source_file=str(Path(path).resolve()),viewing='Excel D65 / 10°',raw=payload)


def _import_chromatic_color_data(path: str, wb) -> ExcelImportResult:
    ws=wb['颜色数据']; headers=_sheet_header_map(ws,1); spectrum=_exported_spectrum_index(wb)
    required={'名称','L*','a*','b*'}
    if not required.issubset(headers):
        raise ValueError('“颜色数据”缺少名称/L*/a*/b*列')
    result=ExcelImportResult(); occupied=set()
    for rn,row in enumerate(ws.iter_rows(min_row=2,values_only=True),2):
        get=lambda key: row[headers[key]] if key in headers and headers[key]<len(row) else None
        name=str(get('名称') or '').strip()
        if not name: continue
        lab=tuple(_as_float(get(k)) for k in ('L*','a*','b*'))
        if any(v is None for v in lab):
            result.warnings.append(f'{name}：Lab 数据不完整，已跳过');continue
        xyz_vals=tuple(_as_float(get(k)) for k in ('X','Y','Z')) if all(k in headers for k in ('X','Y','Z')) else (None,None,None)
        xyz=None if any(v is None for v in xyz_vals) else xyz_vals
        kind=str(get('类型') or 'STD').strip().upper(); kind='BAT' if kind.startswith('BAT') or '批次' in kind else 'STD'
        raw={'EXCEL_ORIGINAL_SOURCE':str(get('来源文件') or '')}
        sm=_sample_from_export_row(path,'颜色数据',rn,name,lab,xyz,kind,spectrum.get(name),raw)
        result.samples.append(sm)
        r=int(_as_float(get('行')) or 0); c=int(_as_float(get('列')) or 0)
        if r>0 and c>0:
            pos=(r-1,c-1)
            if pos in occupied: result.warnings.append(f'{name}：版位 {r},{c} 与其他色样重复')
            occupied.add(pos); result.layout.append({'sample_id':sm.sample_id,'code':short_code(name),'row':r-1,'column':c-1})
    return result


def _import_chromatic_workbench(path: str, wb) -> ExcelImportResult:
    ws=wb['比色工作台']; headers=_sheet_header_map(ws,3); spectrum=_exported_spectrum_index(wb)
    if '名称' not in headers:
        raise ValueError('“比色工作台”缺少“名称”列')
    candidates={}; result=ExcelImportResult()
    for rn,row in enumerate(ws.iter_rows(min_row=4,values_only=True),4):
        get=lambda key: row[headers[key]] if key in headers and headers[key]<len(row) else None
        name=str(get('名称') or '').strip()
        if not name: continue
        role=str(get('角色') or '').strip(); illum=str(get('光源') or '').strip(); obs=str(get('观察者') or '').strip()
        lab_vals=tuple(_as_float(get(k)) for k in ('L*','a*','b*')) if all(k in headers for k in ('L*','a*','b*')) else (None,None,None)
        spec=spectrum.get(name)
        if any(v is None for v in lab_vals):
            if spec:
                from qtx_core.colorimetry import reflectance_to_xyz_lab
                _,lab_vals=reflectance_to_xyz_lab(spec[1],'D65',spec[0],10)
            else:
                result.warnings.append(f'{name}：既没有完整 Lab，也没有光谱，已跳过');continue
        kind='BAT' if '批次' in role else 'STD'
        # Workbench exports can contain one row per illuminant. Prefer the D65/10°
        # record so re-import does not duplicate the same physical measurement.
        score=(0 if illum.upper()=='D65' else 1,0 if '10' in obs else 1,rn)
        key=(name,kind)
        if key not in candidates or score<candidates[key][0]:
            candidates[key]=(score,rn,lab_vals,spec,{'EXCEL_ROLE':role,'EXCEL_ILLUMINANT':illum,'EXCEL_OBSERVER':obs})
    for (name,kind),(_score,rn,lab,spec,raw) in candidates.items():
        result.samples.append(_sample_from_export_row(path,'比色工作台',rn,name,lab,None,kind,spec,raw))
    return result


def _find_generic_header(ws):
    for rn,row in enumerate(ws.iter_rows(min_row=1,max_row=min(20,ws.max_row),values_only=True),1):
        labels={str(v or '').strip():i for i,v in enumerate(row) if str(v or '').strip()}
        name_key=next((k for k in ('名称','色样名称','Name','NAME','Sample Name') if k in labels),None)
        if name_key and all(k in labels for k in ('L*','a*','b*')):
            return rn,labels,name_key
    return None,None,None


def _import_generic_lab_workbook(path: str, wb) -> ExcelImportResult:
    ws=wb[wb.sheetnames[0]]; rn,headers,name_key=_find_generic_header(ws)
    if rn is None:
        raise ValueError('未识别到可导入的数据表。支持：本软件导出的“颜色数据/比色工作台”、旧 MPC 三表，或含 名称 + L*/a*/b* 列的普通 Excel。')
    result=ExcelImportResult()
    for row_number,row in enumerate(ws.iter_rows(min_row=rn+1,values_only=True),rn+1):
        def get(key):
            idx=headers.get(key); return row[idx] if idx is not None and idx<len(row) else None
        name=str(get(name_key) or '').strip()
        if not name: continue
        lab=tuple(_as_float(get(k)) for k in ('L*','a*','b*'))
        if any(v is None for v in lab): continue
        kind_text=str(get('类型') or get('角色') or '')
        kind='BAT' if ('BAT' in kind_text.upper() or '批次' in kind_text) else 'STD'
        result.samples.append(_sample_from_export_row(path,ws.title,row_number,name,lab,None,kind,None,{}))
    if not result.samples: raise ValueError('找到了 Lab 表头，但没有识别到有效色样行')
    result.warnings.append('该 Excel 没有光谱数据；已按 Lab 色样导入，不能凭空生成反射率。')
    return result


def import_excel_workbook(path: str) -> ExcelImportResult:
    """Unified Excel round-trip importer.

    Keeps the legacy MPC three-sheet format, while also accepting workbooks
    exported by Chromatic Analysis itself. This closes the former one-way Excel
    path where the program could export a workbook that it could not open again.
    """
    wb=load_workbook(path,data_only=True,read_only=True)
    sheets=set(wb.sheetnames)
    if {'CIE Coordinates','Reflectance Data','Attributes'}.issubset(sheets):
        wb.close(); return import_mpc_data_workbook(path)
    try:
        if '颜色数据' in sheets: return _import_chromatic_color_data(path,wb)
        if '比色工作台' in sheets: return _import_chromatic_workbench(path,wb)
        return _import_generic_lab_workbook(path,wb)
    finally:
        try: wb.close()
        except Exception: pass


def import_layout_workbook(path: str) -> ExcelImportResult:
    """读取英文排版表；每个page区块按页码与行序保存。"""
    wb = load_workbook(path, data_only=True, read_only=True)
    ws = wb[wb.sheetnames[0]]
    result = ExcelImportResult()
    page = 0
    position = 0
    occupied = set()
    for row in ws.iter_rows(values_only=True):
        first = str(row[0] or "").strip().lower()
        if first.startswith("2025 smart color page"):
            page += 1
            position = 0
            continue
        name = str(row[1] or "").strip() if len(row) > 1 else ""
        code_name = str(row[2] or "").strip() if len(row) > 2 else ""
        code = short_code(code_name)
        if not code.startswith("4P"):
            continue
        key = (page - 1, position)
        if key in occupied:
            result.warnings.append(f"{code}：排版位置冲突 page {page}, row {position + 1}")
        occupied.add(key)
        result.layout.append({"code": code, "commercial_name": name,
                              "page": max(0, page - 1), "row": position, "column": max(0, page - 1)})
        position += 1
    return result


def _lab_to_hex(lab: tuple[float, float, float]) -> str:
    from qtx_core.preview import lab_to_srgb_preview
    return lab_to_srgb_preview(lab).hex[1:]


def _lab_to_rgb(lab: tuple[float, float, float]) -> tuple[int, int, int]:
    from qtx_core.preview import lab_to_srgb_preview
    return lab_to_srgb_preview(lab).rgb8


def _append_sample_metadata_sheets(wb: Workbook, samples: list[Sample], *, include_attributes: bool = True) -> None:
    """Append spectrum, measurement metadata and editable Attributes to an export workbook."""
    samples=[s for s in samples if s is not None]
    if not samples:
        return
    # Reflectance: one row per sample, one column per wavelength in the union grid.
    waves=sorted({int(w) for s in samples for w in (s.wavelengths if s.has_spectrum() else ())})
    rs=wb.create_sheet('光谱数据'); rs.sheet_view.showGridLines=False
    headers=['名称','来源文件',*[f'{w} nm' for w in waves]]
    for c,h in enumerate(headers,1):
        cell=rs.cell(1,c,h); cell.font=Font(bold=True); cell.fill=PatternFill('solid',fgColor='E9EDF2'); cell.alignment=Alignment(horizontal='center')
    for r,s in enumerate(samples,2):
        rs.cell(r,1,s.display_name); rs.cell(r,2,Path(s.source_file).name)
        spec={int(w):float(v) for w,v in zip(s.wavelengths,s.reflectance)} if s.has_spectrum() else {}
        for c,w in enumerate(waves,3):
            if w in spec: rs.cell(r,c,spec[w]).number_format='0.000000'
    rs.freeze_panes='C2'; rs.column_dimensions['A'].width=42; rs.column_dimensions['B'].width=28
    for c in range(3,len(headers)+1):rs.column_dimensions[get_column_letter(c)].width=11

    if include_attributes:
        # Attributes: dynamic fields, not hard-coded to one customer or textile programme.
        attr_keys=[]
        for s in samples:
            for k in (s.raw or {}):
                k=str(k)
                if k.startswith('ATTR_') and k[5:] not in attr_keys: attr_keys.append(k[5:])
        at=wb.create_sheet('属性数据'); at.sheet_view.showGridLines=False
        aheaders=['名称','来源文件',*attr_keys]
        for c,h in enumerate(aheaders,1):
            cell=at.cell(1,c,h); cell.font=Font(bold=True); cell.fill=PatternFill('solid',fgColor='E9EDF2'); cell.alignment=Alignment(horizontal='center')
        for r,s in enumerate(samples,2):
            at.cell(r,1,s.display_name); at.cell(r,2,Path(s.source_file).name); raw=dict(s.raw or {})
            for c,k in enumerate(attr_keys,3):at.cell(r,c,raw.get('ATTR_'+k,''))
        at.freeze_panes='C2'; at.column_dimensions['A'].width=42; at.column_dimensions['B'].width=28
        for c in range(3,len(aheaders)+1):at.column_dimensions[get_column_letter(c)].width=20

    # Measurement summary: preserve the fields needed for traceability when leaving the application.
    ms=wb.create_sheet('测量信息'); ms.sheet_view.showGridLines=False
    mheaders=['名称','来源文件','Viewing','Instrument Model','Instrument Serial','Instrument Settings','Measured/DateTime']
    for c,h in enumerate(mheaders,1):
        cell=ms.cell(1,c,h); cell.font=Font(bold=True); cell.fill=PatternFill('solid',fgColor='E9EDF2'); cell.alignment=Alignment(horizontal='center')
    for r,s in enumerate(samples,2):
        raw=dict(s.raw or {}); prefix='STD' if s.kind=='STD' else 'BAT'
        values=[s.display_name,Path(s.source_file).name,s.viewing,
                raw.get('CPX_INSTRUMENTMODEL') or raw.get('ATTR_Model',''),
                raw.get('CPX_INSTRUMENTSERIAL') or raw.get(f'{prefix}_INSTRUMENT_SERIAL_NO') or raw.get('ATTR_Serial number',''),
                raw.get('CPX_INSTRUMENTSETTINGS') or s.viewing,
                raw.get('CPX_MEASURED') or raw.get(f'{prefix}_DATETIME') or raw.get(f'{prefix}_TIME','')]
        for c,v in enumerate(values,1):ms.cell(r,c,v)
    for c,w in enumerate([42,28,30,22,22,34,24],1):ms.column_dimensions[get_column_letter(c)].width=w

def _save_exchange_workbook(wb: Workbook, path: str) -> None:
    # Exported names, attributes and table strings are literal user data, never
    # executable Excel formulas. Numeric measurement cells stay numeric.
    for sheet in wb:
        for row in sheet:
            for cell in row:
                if cell.data_type == 'f':
                    cell.data_type = 's'
    wb.save(path)


def export_client_card(path: str, title: str, samples: list[Sample | None], columns_count: int = 6,
                       display_hexes: list[str | None] | None = None, app_version: str = "",
                       *, include_internal_metadata: bool = False) -> None:
    """Export the *current palette layout* as a presentation sheet plus data.

    HF86 keeps layout and display colour as single sources of truth: ``samples``
    Customer exports exclude arbitrary internal Attributes by default. Only an
    explicit ``include_internal_metadata=True`` creates the internal attribute sheet.
    ``samples`` contains one entry per current slot (``None`` means an editable blank), and
    optional ``display_hexes`` comes from the same Qt preview conversion used by
    the on-screen cards.  No compaction or re-sorting occurs during export.
    """
    wb=Workbook(); ws=wb.active; ws.title='色卡预览'; ws.sheet_view.showGridLines=False
    columns_count=max(1,min(20,int(columns_count or 6)))
    slots=list(samples); display_hexes=list(display_hexes or [])
    if len(display_hexes)<len(slots):display_hexes.extend([None]*(len(slots)-len(display_hexes)))
    border_side=Side(style='thin',color='D9E1EA'); card_border=Border(left=border_side,right=border_side,top=border_side,bottom=border_side)
    blank_side=Side(style='dashed',color='D9E2EC'); blank_border=Border(left=blank_side,right=blank_side,top=blank_side,bottom=blank_side)
    header_fill=PatternFill('solid',fgColor='F5F7FA')
    # Title block
    ws.merge_cells(start_row=1,start_column=1,end_row=1,end_column=columns_count)
    t=ws.cell(1,1,title); t.font=Font(size=18,bold=True,color='172033'); t.alignment=Alignment(horizontal='left',vertical='center'); ws.row_dimensions[1].height=32
    ws.merge_cells(start_row=2,start_column=1,end_row=2,end_column=columns_count)
    occupied=sum(x is not None for x in slots); blanks=len(slots)-occupied
    sub=ws.cell(2,1,f'{occupied} 色样 · {columns_count} 列 · 空位 {blanks} · D65 / 10°'); sub.font=Font(size=9,color='64748B'); sub.alignment=Alignment(horizontal='left',vertical='center'); ws.row_dimensions[2].height=19
    for col in range(1,columns_count+1):ws.column_dimensions[get_column_letter(col)].width=23

    # One visual card = swatch row + name row + Lab row + slim spacer.
    blocks=max(1,math.ceil(len(slots)/columns_count))
    for index,sample in enumerate(slots):
        col=index%columns_count+1; block=index//columns_count; color_row=4+block*4; name_row=color_row+1; lab_row=color_row+2; spacer_row=color_row+3
        ws.row_dimensions[color_row].height=74; ws.row_dimensions[name_row].height=28; ws.row_dimensions[lab_row].height=19; ws.row_dimensions[spacer_row].height=8
        if sample is None:
            c=ws.cell(color_row,col); c.fill=PatternFill('solid',fgColor='FBFCFE'); c.border=blank_border
            continue
        L,a,b=sample.lab_d65_10; C=math.hypot(a,b); h=math.degrees(math.atan2(b,a))%360
        hex_value=(display_hexes[index] or _lab_to_hex(sample.lab_d65_10)).replace('#','').upper()
        c=ws.cell(color_row,col); c.fill=PatternFill('solid',fgColor=hex_value); c.border=card_border
        n=ws.cell(name_row,col,sample.display_name); n.font=Font(size=9,bold=True,color='20242A'); n.alignment=Alignment(horizontal='left',vertical='center',wrap_text=True); n.fill=PatternFill('solid',fgColor='FFFFFF'); n.border=Border(left=border_side,right=border_side)
        info=ws.cell(lab_row,col,f'L* {L:.1f}   C* {C:.1f}   h° {h:.1f}'); info.font=Font(size=8,color='64748B'); info.alignment=Alignment(horizontal='left',vertical='center'); info.fill=PatternFill('solid',fgColor='FFFFFF'); info.border=Border(left=border_side,right=border_side,bottom=border_side)

    last_row=4+blocks*4
    ws.freeze_panes='A4'; ws.sheet_properties.pageSetUpPr.fitToPage=True; ws.page_setup.orientation='landscape'; ws.page_setup.fitToWidth=1; ws.page_setup.fitToHeight=0
    ws.sheet_properties.pageSetUpPr.autoPageBreaks=False
    ws.page_margins.left=0.25; ws.page_margins.right=0.25; ws.page_margins.top=0.45; ws.page_margins.bottom=0.45
    ws.print_options.horizontalCentered=True
    ws.print_area=f'A1:{get_column_letter(columns_count)}{max(4,last_row)}'
    ws.oddHeader.center.text=title; ws.oddHeader.center.size=10; ws.oddHeader.center.color='64748B'
    ws.oddFooter.left.text='Chromatic Analysis · 色卡工作簿'; ws.oddFooter.right.text='第 &P / &N 页'

    # Full data table keeps explicit layout coordinates instead of flattening the
    # palette into a different order.
    data_ws=wb.create_sheet('颜色数据'); data_ws.sheet_view.showGridLines=False
    headers=['序号','行','列','色块','名称','L*','a*','b*','C*','h°','X','Y','Z','R','G','B','HEX','类型','来源文件']
    for col,header in enumerate(headers,1):
        cell=data_ws.cell(1,col,header); cell.font=Font(bold=True,color='FFFFFF'); cell.alignment=Alignment(horizontal='center',vertical='center'); cell.fill=PatternFill('solid',fgColor='334155')
    data_ws.freeze_panes='A2'; data_ws.auto_filter.ref=f'A1:{get_column_letter(len(headers))}1'; data_ws.row_dimensions[1].height=24
    data_row=2; seq=0
    for index,sample in enumerate(slots):
        if sample is None:continue
        seq+=1; row=index//columns_count+1; col=index%columns_count+1
        L,a,b=sample.lab_d65_10; C=math.hypot(a,b); h=math.degrees(math.atan2(b,a))%360; X,Y,Z=sample.xyz_d65_10
        hex_value=(display_hexes[index] or _lab_to_hex(sample.lab_d65_10)).replace('#','').upper(); R,G,B=tuple(int(hex_value[i:i+2],16) for i in (0,2,4))
        values=[seq,row,col,None,sample.display_name,L,a,b,C,h,X,Y,Z,R,G,B,f'#{hex_value}',sample.kind,Path(sample.source_file).name]
        for cc,value in enumerate(values,1):
            cell=data_ws.cell(data_row,cc,value); cell.alignment=Alignment(horizontal='center' if cc not in (5,19) else 'left',vertical='center',wrap_text=True)
            if 6<=cc<=13:cell.number_format='0.00'
            if data_row%2==0:cell.fill=PatternFill('solid',fgColor='F8FAFC')
        data_ws.cell(data_row,4).fill=PatternFill('solid',fgColor=hex_value); data_ws.row_dimensions[data_row].height=25; data_row+=1
    widths=[8,8,8,12,34,10,10,10,10,10,11,11,11,8,8,8,13,11,28]
    for col,width in enumerate(widths,1):data_ws.column_dimensions[get_column_letter(col)].width=width

    note=wb.create_sheet('导出说明'); note.sheet_view.showGridLines=False
    rows=[('项目','内容'),('方案名称',title),('色样总数',occupied),('空位数量',blanks),('布局',f'{columns_count} 列 × {blocks} 行'),('导出版本',app_version or 'Chromatic Analysis'),('说明','“色卡预览”用于展示与打印；“颜色数据/光谱数据/测量信息”保留可分析的数据；默认不导出内部属性。色块为屏幕预览，光谱反射率及仪器测量值仍为权威记录。')]
    for r,(k,v) in enumerate(rows,1):
        note.cell(r,1,k); note.cell(r,2,v)
        note.cell(r,1).font=Font(bold=True,color='334155'); note.cell(r,1).fill=header_fill
        note.cell(r,1).alignment=Alignment(vertical='top'); note.cell(r,2).alignment=Alignment(vertical='top',wrap_text=True)
    note.column_dimensions['A'].width=18; note.column_dimensions['B'].width=70
    _append_sample_metadata_sheets(wb,[x for x in slots if x is not None], include_attributes=include_internal_metadata); _save_exchange_workbook(wb, path)



def export_workbench_table(path: str, title: str, headers: list[str], rows: list[list], samples: list[Sample] | None = None) -> None:
    """按工作台当前可见列原样导出；色块列写入真实填充色。"""
    wb = Workbook()
    ws = wb.active
    ws.title = "比色工作台"
    ws.sheet_view.showGridLines = False
    count = max(1, len(headers))
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=count)
    ws.cell(1, 1, title).font = Font(size=16, bold=True)
    ws.cell(1, 1).alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 28
    for c, header in enumerate(headers, 1):
        cell = ws.cell(3, c, header)
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.fill = PatternFill("solid", fgColor="E9EDF2")
    ws.freeze_panes = "A4"
    ws.auto_filter.ref = f"A3:{get_column_letter(count)}3"
    for r_idx, row in enumerate(rows, 4):
        ws.row_dimensions[r_idx].height = 30
        for c_idx, value in enumerate(row, 1):
            cell = ws.cell(r_idx, c_idx)
            if isinstance(value, dict) and "swatch_lab" in value:
                cell.fill = PatternFill("solid", fgColor=_lab_to_hex(tuple(value["swatch_lab"])))
                cell.value = ""
            else:
                cell.value = value
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for c, header in enumerate(headers, 1):
        width = 12
        if header == "色块": width = 14
        elif header == "名称": width = 28
        elif header in {"判定", "角色"}: width = 14
        elif "ΔE" in header or header in {"CMC", "CIE94", "MI", "WI", "Tint"}: width = 12
        ws.column_dimensions[get_column_letter(c)].width = width
    if samples:
        _append_sample_metadata_sheets(wb,list(samples))
    _save_exchange_workbook(wb, path)
