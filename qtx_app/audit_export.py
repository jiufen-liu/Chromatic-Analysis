from __future__ import annotations

import time
from pathlib import Path
from typing import Iterable, Mapping, Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


_HEADER_FILL = PatternFill('solid', fgColor='EAF2FF')
_TITLE_FILL = PatternFill('solid', fgColor='F7F9FC')
_THIN = Side(style='thin', color='D8DEE9')
_BORDER = Border(bottom=_THIN)


def _value(row: Mapping[str, Any], key: str, default: Any = '') -> Any:
    try:
        return row[key]
    except Exception:
        return default


def export_audit_workbook(
    path: str | Path,
    rows: Iterable[Mapping[str, Any]],
    *, organization_name: str,
    installation_id: str,
    installation_label: str,
    exported_by: str,
    filter_description: str,
) -> Path:
    """Export filtered audit rows as a two-sheet administrator workbook."""
    target = Path(path)
    if target.suffix.lower() != '.xlsx':
        target = target.with_suffix('.xlsx')
    target.parent.mkdir(parents=True, exist_ok=True)

    records = list(rows)
    wb = Workbook()
    ws = wb.active
    ws.title = '审计记录'
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = 'A2'

    headers = ['序号', '时间', '用户', '操作类型', '目标对象', '详细信息']
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True, color='1F2937')
        cell.fill = _HEADER_FILL
        cell.alignment = Alignment(horizontal='center', vertical='center')
        cell.border = _BORDER
    ws.row_dimensions[1].height = 24

    for index, row in enumerate(records, 1):
        created = float(_value(row, 'created_at', 0) or 0)
        timestamp = time.strftime('%Y-%m-%d %H:%M:%S', time.localtime(created)) if created else ''
        ws.append([
            index,
            timestamp,
            str(_value(row, 'username', '') or ''),
            str(_value(row, 'action', '') or ''),
            str(_value(row, 'target', '') or ''),
            str(_value(row, 'detail', '') or ''),
        ])
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical='top', wrap_text=True)
            cell.border = _BORDER
    ws.auto_filter.ref = f'A1:F{max(1, ws.max_row)}'
    widths = [9, 21, 18, 30, 42, 52]
    for idx, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(idx)].width = width

    info = wb.create_sheet('导出说明')
    info.sheet_view.showGridLines = False
    info['A1'] = 'Chromatic Analysis 审计记录导出说明'
    info['A1'].font = Font(size=14, bold=True, color='1F2937')
    info['A1'].fill = _TITLE_FILL
    info.merge_cells('A1:B1')
    meta = [
        ('组织', organization_name),
        ('Installation ID', installation_id),
        ('本机实例', installation_label),
        ('导出人', exported_by),
        ('导出时间', time.strftime('%Y-%m-%d %H:%M:%S')),
        ('筛选条件', filter_description or '全部记录'),
        ('记录总数', len(records)),
        ('说明', '此文件由本机组织管理员导出。导出行为本身会写入审计日志。'),
    ]
    for row_index, (key, value) in enumerate(meta, 3):
        info.cell(row_index, 1, key)
        info.cell(row_index, 2, value)
        info.cell(row_index, 1).font = Font(bold=True, color='475467')
        info.cell(row_index, 1).fill = _TITLE_FILL
        info.cell(row_index, 1).alignment = Alignment(vertical='top')
        info.cell(row_index, 2).alignment = Alignment(vertical='top', wrap_text=True)
        info.cell(row_index, 1).border = info.cell(row_index, 2).border = _BORDER
    info.column_dimensions['A'].width = 22
    info.column_dimensions['B'].width = 76

    wb.save(target)
    return target
